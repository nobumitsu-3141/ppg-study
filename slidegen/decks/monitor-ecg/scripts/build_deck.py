#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""content/*.json（Opus執筆の医学中身）＋ figs/*.png → 川副式pptx。
編集して再実行すれば作り直せる。"""
import os, json, glob
from decklib import *
from decklib import _txbox, _set_font, _set_title, _strip_body
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

HERE = os.path.dirname(os.path.abspath(__file__))
FIGDIR = os.path.abspath(os.path.join(HERE, "..", "figs"))
CONTENT = os.path.abspath(os.path.join(HERE, "..", "content"))
OUT = os.path.abspath(os.path.join(HERE, "..", "モニター心電図.pptx"))
def F(n): return os.path.join(FIGDIR, n + ".png")

TEMPLATE = os.path.join(HERE, "_template_kawazoe.pptx")

CMAP = {"INK": INK, "GOLD": GOLD, "RED": RED, "GRAY": GRAY, "TEAL": TEAL,
        "BLUE": BLUE, "GREEN": RGBColor(0x54,0x82,0x35), "MGRAY": MGRAY}
ACC = {"TEAL": TEAL, "BLUE": BLUE, "GOLD": GOLD}

# code → figure file (コードで対応：エージェントの figure.id には依存しない)
FIGMAP = {
 "1.1":"fig_1_1_prehistory","1.2":"fig_1_2_einthoven","1.3":"fig_1_3_twelve_lead","1.4":"fig_1_4_monitoring",
 "2.1":"fig_2_1_action_potential","2.2":"fig_2_2_conduction","2.3":"fig_2_3_vector","2.4":"fig_2_4_projection",
 "2.5":"fig_2_5_lead_ii","2.6":"fig_2_6_signal_chain",
 "3.1":"fig_3_1_three_lead","3.2":"fig_3_2_five_lead","3.4":"fig_3_4_complex","3.5":"fig_3_5_intervals",
 "3.6":"fig_3_6_rate","3.7":"fig_3_7_axis",
 "4.1":"fig_4_1_nsr","4.2":"fig_4_2_normal_values","4.3":"fig_4_3_normal_stt","4.4":"fig_4_4_monitor_normal",
 "5.2":"fig_5_2_leads","5.3":"fig_5_3_st","5.5":"fig_5_5_artifact","5.7":"fig_5_7_rhythm",
 "6.1":"fig_6_1_read_order","6.2":"fig_6_2_brady","6.3":"fig_6_3_narrow","6.4":"fig_6_4_wide",
 "6.5":"fig_6_5_pac_pvc","6.6":"fig_6_6_lethal","6.7":"fig_6_7_bbb","6.8":"fig_6_8_ischemia","6.9":"fig_6_9_electrolyte",
 "7.1":"fig_7_1_principle","7.3":"fig_7_3_tachy","7.4":"fig_7_4_vf",
 "8.1":"fig_8_1_pacemaker","8.3":"fig_8_3_transplant","8.5":"fig_8_5_peds","8.6":"fig_8_6_special",
}

FULL = ["心電図モニタリングの歴史",
        "信号の成り立ち（活動電位→体表）",
        "モニター心電図の基礎",
        "正常波形の見方",
        "麻酔中の見どころ",
        "異常所見の判読",
        "異常への対応",
        "特殊な使用例"]
SHORT = ["歴史", "成り立ち", "基礎", "正常", "術中の見どころ", "異常所見", "対応", "特殊な使用例"]
SUBS = [
 [("1.1–1.4", "前史・Einthoven・12誘導・連続監視"), ("1.5", "麻酔とモニタ心電図")],
 [("2.1–2.3", "活動電位・刺激伝導系・電気ベクトル"), ("2.4–2.6", "誘導軸・なぜII誘導・信号処理")],
 [("3.1–3.3", "電極配置(3点/5点)・変法誘導"), ("3.4–3.7", "波形の構成・間隔・心拍数・電気軸")],
 [("4.1–4.2", "正常洞調律・各波の正常値"), ("4.3–4.4", "正常ST-T・モニタでの正常像")],
 [("5.1–5.4", "監視項目・誘導選択・ST解析・フィルタ"), ("5.5–5.7", "アーチファクト・設定・術中の調律変化")],
 [("6.1–6.6", "系統判読・徐脈/頻脈・期外・致死的"), ("6.7–6.10", "脚ブロック・虚血・電解質・鑑別")],
 [("7.1–7.4", "原則・徐脈・頻脈・VF/無脈性VT"), ("7.5–7.7", "PEA/心静止・術中虚血・高K/TdP/LAST")],
 [("8.1–8.3", "ペースメーカ・CIED・移植心"), ("8.4–8.6", "特殊電極・小児・特殊病態の露見")],
]
MENUSUB = ["前史 → Einthoven → 12誘導 → 連続監視 → 麻酔",
           "活動電位・伝導系・ベクトル・誘導軸・信号処理",
           "電極配置・波形と間隔・心拍数・電気軸",
           "正常洞調律・正常値・正常ST-T・モニタ像",
           "監視項目・II+V5・ST自動解析・アーチ・術中変化",
           "系統判読・不整脈・脚ブロック・虚血・電解質",
           "原則・徐脈/頻脈・VF・PEA・虚血・高K/LAST",
           "ペースメーカ・CIED・移植心・小児・特殊病態"]

# ---------------- load content ----------------
def load_content():
    slides = []
    for fn in ["ch1_2.json", "ch3_4.json", "ch5.json", "ch6.json", "ch7_8.json"]:
        with open(os.path.join(CONTENT, fn), encoding="utf-8") as f:
            slides += json.load(f)
    slides.sort(key=lambda s: [int(x) for x in s["code"].split(".")])
    return slides

# ---------------- factories ----------------
prs = new_deck(TEMPLATE)

def divider(idx):
    s = content_slide(prs, f"{idx+1}. {SHORT[idx]}", active=idx)
    tb, tf = _txbox(s, Cm(2.4), Cm(4.8), Cm(29.4), Cm(14.0))
    for i, name in enumerate(FULL):
        on = (i == idx)
        add_para(tf, [(f"{i+1}. ", dict(size=26 if on else 22, bold=True, color=GOLD if on else LGRAY)),
                      (name, dict(size=26 if on else 22, bold=on, color=INK if on else LGRAY))],
                 first=(i == 0), space_after=(5 if on else 9), line=1.05)
        if on:
            for code, label in SUBS[i]:
                add_para(tf, [(code + "　", dict(size=22, bold=True, color=GOLD)),
                              (label, dict(size=22, color=INK))], level=1, space_after=6, line=1.05)
            tf.paragraphs[-1].space_after = Pt(12)
    return s

def bullets(title, items, active=None, cite=None, cite_url=None, size=22, lead=None):
    s = content_slide(prs, title, active=active, cite=cite, cite_url=cite_url)
    y = Cm(5.6)
    if lead:
        tb, tf = _txbox(s, Cm(2.4), y, Cm(29.0), Cm(1.4))
        add_para(tf, [(lead, dict(size=18, bold=True, color=GOLD))], first=True)
        y = Cm(7.2)
    tb, tf = _txbox(s, Cm(2.6), y, Cm(28.6), Cm(12.0))
    for i, (txt, kw) in enumerate(items):
        kw = dict(kw); sa = kw.pop("sa", 16); lv = kw.pop("lv", 0); kw.setdefault("size", size)
        add_para(tf, [(txt, kw)], first=(i == 0), space_after=sa, level=lv, line=1.12)
    return s

def column(title, heading, body, active=None, cite=None, cite_url=None, accent=TEAL):
    s = content_slide(prs, title, active=active, cite=cite, cite_url=cite_url)
    add_box(s, Cm(2.2), Cm(5.1), Cm(29.5), Cm(12.4), RGBColor(0xF5, 0xF7, 0xF9), line=accent)
    tb, tf = _txbox(s, Cm(3.0), Cm(5.8), Cm(27.9), Cm(11.2))
    add_para(tf, [(heading, dict(size=26, bold=True, color=accent))], first=True, space_after=16, line=1.05)
    for para in body:
        add_para(tf, [("●  ", dict(size=22, color=accent)), (para, dict(size=22, color=INK))],
                 space_after=14, line=1.2)
    return s

def grid2(title, cells, active=None, cite=None):
    s = content_slide(prs, title, active=active, cite=cite)
    cols, rows = 2, 3
    x0, y0 = Cm(2.4), Cm(5.4)
    cw, ch = Cm(14.5), Cm(3.9); gx, gy = Cm(0.7), Cm(0.35)
    for i, cell in enumerate(cells[:6]):
        r, c = divmod(i, cols)
        x = x0 + c * (cw + gx); y = y0 + r * (ch + gy)
        add_box(s, x, y, cw, ch, GOLDBG if False else RGBColor(0xF5,0xF7,0xF9), line=TEAL)
        tb, tf = _txbox(s, x + Cm(0.5), y + Cm(0.35), cw - Cm(1.0), ch - Cm(0.7), anchor=MSO_ANCHOR.MIDDLE)
        add_para(tf, [(f"{i+1}. ", dict(size=22, bold=True, color=GOLD)),
                      (cell["h"], dict(size=22, bold=True, color=INK))], first=True, space_after=4, line=1.05)
        add_para(tf, [("　" + cell["d"], dict(size=22, color=GRAY))], space_after=0, line=1.0)
    return s

def references(refs, title="参考文献"):
    s = content_slide(prs, title, active=None)
    col_w = Cm(14.6); half = (len(refs) + 1) // 2
    for c, (x, chunk) in enumerate([(Cm(2.2), refs[:half]), (Cm(17.4), refs[half:])]):
        tb, tf = _txbox(s, x, Cm(5.1), col_w, Cm(13.4))
        for i, (txt, url) in enumerate(chunk):
            p = add_para(tf, [(txt, dict(size=10.5, color=GRAY))], first=(i == 0), space_after=7, line=1.05)
            if url and p.runs:
                p.runs[-1].hyperlink.address = url
                p.runs[-1].font.underline = False
                p.runs[-1].font.color.rgb = GRAY
    return s

def render_slide(sl):
    code = sl["code"]; ch = sl["chapter"]; active = ch - 1
    title = f"{code}　{sl['title']}"
    typ = sl["type"]
    cite = sl.get("cite"); url = sl.get("cite_url")
    if typ in ("figure", "table"):
        fig = FIGMAP.get(code)
        if not fig or not os.path.exists(F(fig)):
            print("!! missing fig for", code, fig); return
        figure_slide(prs, title, F(fig), active=active, cite=cite, cite_url=url,
                     caption=sl.get("caption"))
    elif typ == "bullets":
        items = [(b["text"], dict(color=CMAP.get(b.get("color", "INK"), INK))) for b in sl["bullets"]]
        s = bullets(title, items, active=active, cite=cite, cite_url=url)
        set_notes(s, sl.get("notes"))
    elif typ == "column":
        col = sl["column"]
        s = column(title, col["heading"], col["body"], active=active, cite=cite, cite_url=url,
                   accent=ACC.get(col.get("accent", "TEAL"), TEAL))
        set_notes(s, sl.get("notes"))
    elif typ == "grid2":
        grid2(title, sl["grid"], active=active, cite=cite)
    else:
        print("!! unknown type", typ, code)

# ============================ BUILD ============================
slides = load_content()

# 1 title
title_slide(prs, "モニター心電図の読み方",
            [("― 成り立ちから術中対応まで ―", dict(size=22, bold=True, color=INK)),
             ("麻酔科専門医向け 勉強会", dict(size=20, bold=False, color=GRAY)),
             ("", dict(size=10, color=INK)),
             ("川副 靖晃", dict(size=22, bold=False, color=INK))],
            hero=F("fig_hero"))

# 2 clinical question
bullets("モニター心電図、読めていますか？",
        [("毎日見ている II・V5 の一本の線。その形とトレンドから何を読むのか？", dict(color=INK, sa=18)),
         ("術中に見逃してはいけない 虚血・不整脈・電解質・QT のサインは？", dict(color=INK, sa=18)),
         ("そして波形が乱れたとき、まず何を見て、どう動くのか？", dict(color=INK, sa=24)),
         ("→ 成り立ちから対応まで、系統立てて整理する", dict(color=RED, sa=0))],
        size=22)

# 3 goals
s = bullets("目標", [
    ("① 心電図の成り立ち（活動電位→ベクトル→誘導→波形）を理解する", dict(color=INK, sa=16)),
    ("② 正常と異常を読む（レート→調律→P→PR→QRS→ST/T→QT）", dict(color=INK, sa=16)),
    ("③ 術中の見どころ（II+V5 の虚血監視・アーチファクト・QT）を掴む", dict(color=INK, sa=16)),
    ("④ 代表的な異常への 初期対応 を、ガイドライン準拠で押さえる", dict(color=INK, sa=0)),
], size=22)
set_notes(s, "目標\n"
              "① 心電図信号の 成り立ち（活動電位→ベクトル→誘導→波形）を理解する\n"
              "② 正常と異常を 系統的に 読む（レート→調律→P→PR→QRS→ST/T→QT）\n"
              "③ 術中の見どころ（II+V5 の虚血監視・アーチファクト・QT）を掴む\n"
              "④ 代表的な異常への 初期対応 を、ガイドライン準拠で押さえる")

# 4 menu（各行1行厳守。章ごとの小見出しはスピーカーノートへ）
s = content_slide(prs, "メニュー", active=-1)
tb, tf = _txbox(s, Cm(2.4), Cm(4.7), Cm(29.2), Cm(13.9), anchor=MSO_ANCHOR.MIDDLE)
for i, name in enumerate(FULL):
    add_para(tf, [(f"{i+1}.  ", dict(size=24, bold=True, color=GOLD)),
                  (name, dict(size=24, bold=True, color=INK))], first=(i == 0), space_after=10, line=1.0)
tf.paragraphs[-1].space_after = Pt(0)
set_notes(s, "各章の内訳:\n" + "\n".join(f"{i+1}. {FULL[i]} — {MENUSUB[i]}" for i in range(8)))

# chapters
by_ch = {}
for sl in slides:
    by_ch.setdefault(sl["chapter"], []).append(sl)
for ch in range(1, 9):
    divider(ch - 1)
    for sl in by_ch.get(ch, []):
        render_slide(sl)

# references (unique)
seen = set(); refs = []
for sl in slides:
    c = sl.get("cite"); u = sl.get("cite_url")
    if c and c not in seen:
        seen.add(c); refs.append((f"[{len(refs)+1}] {c}", u))
half = (len(refs) + 1) // 2 if len(refs) <= 24 else (len(refs) + 1) // 2
if len(refs) <= 24:
    references(refs)
else:
    references(refs[:24], "参考文献 (1)")
    references(refs[24:], "参考文献 (2)")

add_pagenumbers(prs)
prs.save(OUT)
print("saved", OUT, "slides:", len(prs.slides._sldIdLst))
