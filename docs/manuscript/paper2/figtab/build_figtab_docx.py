#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""論文2 図表集 ― 図 7 点（凡例文つき）と表 18 点を A・B・C の 3 部にまとめた docx を作る。

和文 `論文2_図表集_ja.docx`・英文 `paper2_figures_tables_en.docx`。図は `out/<名前>_<言語>.png`（600 dpi）を
300 dpi に落として図の実寸の幅（150 mm 以下）で貼り、凡例文は `out/<名前>_legend_<言語>.txt` をそのまま段落にする。
表は `build_tables.py` の表（`build_all`）を同じ描き方（`write_docx` と同じ書式）で入れる。数値はどこにも手で打たない。

使い方
    python3 build_figtab_docx.py             両方の docx を書く
    python3 build_figtab_docx.py --zip       あわせて配布用の zip（paper2_figtab.zip）を作る
    python3 build_figtab_docx.py --selftest  図 7・表 18 が入り、図の幅が 150 mm 以下で、凡例文が揃い、禁止語が無いことを確かめる
"""
from __future__ import annotations

import argparse
import io
import re
import sys
import tempfile
import zipfile
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
import build_tables as bt          # noqa: E402
import common                      # noqa: E402

OUT = HERE / "out"
DOCX = {"ja": HERE / "論文2_図表集_ja.docx", "en": HERE / "paper2_figures_tables_en.docx"}
ZIP = HERE / "paper2_figtab.zip"
FIG_WIDTH_MM = 150.0

# 3 部の構成。図の名前は out/ のファイル名、表の名前は build_tables の tid
PARTS = [
    {"key": "A",
     "title": {"ja": "A　事前に決めた解析", "en": "Part A. Prespecified analyses"},
     "lead": {"ja": "何を、どの規準で、いつ凍結し、どうなったか。判定は `docs/research/roadmap_v1.md` §9 のまま。",
              "en": "What was tested, under which frozen criterion, and what came out. Verdicts are those of "
                    "`docs/research/roadmap_v1.md` §9."},
     "figures": ["fig_judgement", "fig_effects"],
     "tables": ["table1", "table2", "table3", "table3b", "table3c", "table4", "table5a", "table5b"]},
    {"key": "B",
     "title": {"ja": "B　解析を踏まえた改善の設計", "en": "Part B. Design of the improvements"},
     "lead": {"ja": "何が分かり、何を変えようとし、走らせる前に何を予測したか。表4 以降は探索・事後で、判定には用いない。",
              "en": "What the analyses showed, what was changed, and what was predicted before each run. "
                    "Table 4 onward is exploratory and post hoc and is not used for the decision."},
     "figures": ["fig_timeline", "fig_synthetic"],
     "tables": ["table6a", "table6b", "table7"]},
    {"key": "C",
     "title": {"ja": "C　改善で結果は変わったか", "en": "Part C. Did the improvements change the result?"},
     "lead": {"ja": "改良案ごとに A・B・C 段で何が変わり、雑音に残るか。すべて探索・事後で、判定には用いない。",
              "en": "What changed per variant in tiers A, B and C, and what survives noise. All exploratory and "
                    "post hoc; not used for the decision."},
     "figures": ["fig_variants", "fig_noise", "fig_tradeoff"],
     "tables": ["table8", "table8b", "table9", "tableS1", "tableS2a", "tableS2b", "tableS3"]},
]
FIG_NO = {"fig_judgement": {"ja": "図1", "en": "Figure 1"}, "fig_effects": {"ja": "図2", "en": "Figure 2"},
          "fig_timeline": {"ja": "図3", "en": "Figure 3"}, "fig_synthetic": {"ja": "図4", "en": "Figure 4"},
          "fig_variants": {"ja": "図5", "en": "Figure 5"}, "fig_noise": {"ja": "図6", "en": "Figure 6"},
          "fig_tradeoff": {"ja": "図7", "en": "Figure 7"}}

TITLE = {"ja": "論文2 図表集 ― 事前に決めた解析 → 改善の設計 → 改善後の結果",
         "en": "Paper 2 figures and tables: prespecified analyses, design of the improvements, results after them"}
INTRO = {
    "ja": [
        "論文2（PWDB 4,374 名の仮想被験者。凍結版の脈波分解は不成立、同じ波形の特徴点法は成立）の結果を、"
        "A 事前に決めた解析、B 解析を踏まえた改善の設計、C 改善で結果は変わったか、の 3 部に分けて図 {n_fig} 点と表 {n_tab} 点にした。",
        "番号はこの図表集の中の仮の番号である（原稿 v2 は図 3 点・表 4 点で、対応は未定）。"
        "数値はすべて `../02_tables.md` を機械で読んだ `data/paper2_numbers.json` から各台本が実行時に読み、手で打っていない"
        "（図4 の合成脈波の値だけは例示で、図の中にその旨を書いた）。",
        "判定（成立・不成立）は `docs/research/roadmap_v1.md` §9 のまま動かさない。表4 以降と図3〜図7 に含まれる改良版・型別・"
        "B 段・雑音・波分離の内容は、決定試験の判定の後に回した探索・事後の記述で、判定には用いない。"
        "改良した当てはめを主解析に用いるには新たな事前登録が要る。",
        "段の注釈: " + common.STAGE_NOTE["ja"] + "。A だけ良くて C が悪ければ、難しい拍を捨てたから良く見えたのであって改善ではない。",
    ],
    "en": [
        "Results of paper 2 (Pulse Wave Database, 4,374 virtual subjects; the frozen pulse decomposition failed the "
        "prespecified criterion, fiducial-point analysis of the same waveforms passed) arranged in three parts: "
        "A, the prespecified analyses; B, how the analyses led to the design of the improvements; C, whether the "
        "improvements changed the result. {n_fig} figures and {n_tab} tables.",
        "Numbering is provisional to this set (manuscript v2 has three figures and four tables; the mapping is to be "
        "decided). Every number is read at run time from `data/paper2_numbers.json`, a mechanical parse of "
        "`../02_tables.md`; none is typed by hand (the only exception is the synthetic-beat illustration of figure 4, "
        "which says so in the figure).",
        "Verdicts (pass or fail) are those of `docs/research/roadmap_v1.md` §9 and are not changed. The contents of "
        "table 4 onward and of figures 3 to 7 (improved fits, waveform types, tier B, noise, wave separation) are "
        "exploratory, post hoc descriptions made after the prespecified decision and are not used for it. Using an "
        "improved fit as a primary analysis would require a new preregistration.",
        "Tier note: " + common.STAGE_NOTE["en"] + ". A gain in tier A alone with no gain in tier C means that "
        "difficult beats were discarded, not that the index improved.",
    ],
}
SOURCES = {
    "ja": [
        "出典と作り直し",
        "図は `build_fig_<名前>.py`、表は `build_tables.py`、この文書は `build_figtab_docx.py` が作る（`build_all.sh` で一括）。"
        "各台本は `--selftest` を持ち、描いた値と JSON の一致、最小文字サイズ 8 pt、段の注釈と探索・事後の注記、文字の重なり、"
        "禁止語（`analysis/scripts/check_terminology.py` の表）を確かめる。",
        "数値の出どころは `../02_tables.md`（表1〜表7b）で、その記録は `docs/research/lab_log.md` の追記12（決定試験）・13（基底）・"
        "18（文献の条件）・137（型別）・139（0.65T の打ち切りを候補 (3) に挙げた）・141（合成脈波の診断）・"
        "143〜147（再当てはめ 5・7・10 型）・149〜153（14 型・B 段の設計）・152（14 型と線形分離）・158（B 段と雑音）・"
        "162（この図表集の点検で直したことと、0.65 の出どころの記述の訂正）にある。事前指定の項目と改善の経過は `data/prespec_chronology.json`・"
        "`notes_prespec_and_chronology.md`。",
        "Mac 1 の結果ファイル（`50_reservoir_bench_C14.txt`・`50_reservoir_bench_B3.txt`・`50_reservoir_bench_noise0.01.txt`・"
        "`50_reservoir_bench_noise0.02.txt`・`51_wave_separation.txt`）はまだリポジトリに無く、該当する数値は `02_tables.md` の記録による。",
        "投稿先（Physiological Measurement, IOP 共通の規定）: 図幅 85 mm（単段）／150 mm（二段）、図の文字は最終寸法で 8 pt 以上、"
        "PDF（ベクタ）・SVG・PNG（600 dpi）。文字は編集可能なテキストとして埋め込んである。CMYK への変換は投稿時に行う。",
    ],
    "en": [
        "Sources and reproduction",
        "Figures are produced by `build_fig_<name>.py`, tables by `build_tables.py`, and this document by "
        "`build_figtab_docx.py` (`build_all.sh` runs everything). Every script has a `--selftest` that checks the plotted "
        "values against the JSON, the 8 pt minimum font size, the tier and post-hoc notes, text overlaps and banned terms "
        "(the list of `analysis/scripts/check_terminology.py`).",
        "The numbers come from `../02_tables.md` (tables 1 to 7b), recorded in `docs/research/lab_log.md` entries 12 "
        "(decision test), 13 (basis functions), 18 (published conditions), 137 (waveform types), 139 (0.65T truncation "
        "listed as candidate (3)), 141 (synthetic-beat diagnosis), 143 to 147 (refits with 5, 7 and 10 variants), 149 to "
        "153 (14 variants and the design of tier B), 152 (14 variants and wave separation), 158 (tier B and noise) and 162 "
        "(corrections made while checking this set, including the stated source of the value 0.65). The prespecified items and the chronology are "
        "in `data/prespec_chronology.json` and `notes_prespec_and_chronology.md`.",
        "The result files written on Mac 1 (`50_reservoir_bench_C14.txt`, `50_reservoir_bench_B3.txt`, "
        "`50_reservoir_bench_noise0.01.txt`, `50_reservoir_bench_noise0.02.txt`, `51_wave_separation.txt`) are not yet in "
        "the repository; the corresponding numbers are those recorded in `02_tables.md`.",
        "Journal specification applied (Physiological Measurement; IOP-wide rules): figure width 85 mm (single column) "
        "or 150 mm (double column), figure text at least 8 pt at final size, vector PDF plus SVG and 600 dpi PNG, text "
        "embedded as editable text. Conversion to CMYK is left to submission.",
    ],
}


_PROVENANCE_LINE = re.compile(r"^\s{2,}\S.*(←|<-)")


def legend_paragraphs(name: str, lang: str) -> list[str]:
    """凡例文のファイルを段落に分ける（空行で区切る。逆引用符は外す）。

    `  名前 = 値 ← JSON の場所` の形の出どころの一覧（作り直しのための記録）は、原稿の凡例文には入れない
    （`out/<名前>_legend_<言語>.txt` にはそのまま残る）。"""
    txt = (OUT / f"{name}_legend_{lang}.txt").read_text(encoding="utf-8")
    paras = []
    for block in re.split(r"\n\s*\n", txt):
        if not block.strip():
            continue
        lines = block.strip("\n").split("\n")
        if sum(bool(_PROVENANCE_LINE.match(ln)) for ln in lines) >= 2:
            continue
        p = block.strip().replace("`", "")
        paras.append(re.sub(r"\s*\n\s*", " " if lang == "en" else "", p))
    return paras


def png_300dpi(name: str, lang: str, tmp: Path) -> Path:
    """600 dpi の PNG を 300 dpi に落として docx に貼る（提出用の 600 dpi はそのまま out/ に残す）。"""
    from PIL import Image
    src = OUT / f"{name}_{lang}.png"
    dst = tmp / f"{name}_{lang}_300.png"
    with Image.open(src) as im:
        w, h = im.size
        im2 = im.resize((w // 2, h // 2), Image.LANCZOS)
        im2.save(dst, dpi=(300, 300))
    return dst


def write_combined(lang: str, path: Path, tables: list) -> None:
    from docx import Document
    from docx.enum.section import WD_ORIENT, WD_SECTION
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Mm, Pt

    font_name = "游明朝" if lang == "ja" else "Times New Roman"
    body_pt = 10.5 if lang == "ja" else 10.0
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = font_name
    st.font.size = Pt(body_pt)
    rpr = st.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), font_name)

    def set_page(sec, landscape: bool):
        if landscape:
            sec.orientation = WD_ORIENT.LANDSCAPE
            sec.page_width, sec.page_height = Mm(297), Mm(210)
        else:
            sec.orientation = WD_ORIENT.PORTRAIT
            sec.page_width, sec.page_height = Mm(210), Mm(297)
        sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Mm(20)

    def style_run(run, size, bold=None, sup_=False):
        run.font.name = font_name
        run.font.size = Pt(size)
        r = run._element.get_or_add_rPr()
        rf = r.find(qn("w:rFonts"))
        if rf is None:
            rf = OxmlElement("w:rFonts")
            r.append(rf)
        rf.set(qn("w:eastAsia"), font_name)
        if bold is not None:
            run.bold = bold
        if sup_:
            run.font.superscript = True

    def add_marked(par, text: str, size: float, base_bold: bool = False):
        text = text.replace("`", "").replace("\\|", "|")
        bold = base_bold
        for tok in bt._TOKEN_RE.split(text):
            if not tok:
                continue
            if tok == "**":
                bold = not bold
                continue
            m = re.fullmatch(r"<sup>([a-z](?:, [a-z])*)</sup>", tok)
            if m:
                style_run(par.add_run(m.group(1)), size, bold=False, sup_=True)
                continue
            style_run(par.add_run(tok), size, bold=bold)

    def heading(text: str, size: float):
        p = doc.add_paragraph()
        style_run(p.add_run(text), size, bold=True)
        return p

    def ensure_orientation(landscape: bool):
        nonlocal cur_landscape
        if landscape != cur_landscape:
            set_page(doc.add_section(WD_SECTION.NEW_PAGE), landscape)
            cur_landscape = landscape
        else:
            doc.add_page_break()

    by_tid = {T.tid: T for T in tables}
    set_page(doc.sections[0], False)
    cur_landscape = False
    heading(TITLE[lang], body_pt + 3)
    n_fig = sum(len(p["figures"]) for p in PARTS)
    n_tab = sum(len(p["tables"]) for p in PARTS)
    for para in INTRO[lang]:
        add_marked(doc.add_paragraph(), para.replace("{n_fig}", str(n_fig)).replace("{n_tab}", str(n_tab)), body_pt)

    with tempfile.TemporaryDirectory() as td:
        tmp = Path(td)
        for part in PARTS:
            ensure_orientation(False)
            heading(part["title"][lang], body_pt + 2)
            add_marked(doc.add_paragraph(), part["lead"][lang], body_pt)
            for k, name in enumerate(part["figures"]):
                # 2 つ目以降の図は新しいページから（最初の図は部の見出しと同じページ）。表は ensure_orientation が
                # 新しいページを始めるので、図の後ろには改ページを置かない（白紙のページができないように）
                if k > 0:
                    doc.add_page_break()
                doc.add_paragraph()
                # 図の実寸の幅で貼る（拡大すると 8 pt の文字が大きくなり、縮めると 8 pt を下回る）
                w_mm = min(common.png_size_mm(OUT / f"{name}_{lang}.png")[0], FIG_WIDTH_MM)
                doc.add_picture(str(png_300dpi(name, lang, tmp)), width=Mm(w_mm))
                paras = legend_paragraphs(name, lang)
                cap = doc.add_paragraph()
                style_run(cap.add_run(paras[0]), body_pt, bold=True)
                for p in paras[1:]:
                    add_marked(doc.add_paragraph(), p, body_pt - 0.5)
            for tid in part["tables"]:
                T = by_tid[tid]
                ensure_orientation(T.landscape)
                cell_pt = 9.0 if T.small else body_pt
                cap = doc.add_paragraph()
                style_run(cap.add_run(f"{T.number[lang]}. "), body_pt, bold=True)
                add_marked(cap, T.mark(T.caption[lang]), body_pt)
                ncols = T.ncols(lang)
                tbl = doc.add_table(rows=1, cols=ncols)
                tbl.style = "Table Grid"
                tbl.autofit = True
                hdr = tbl.rows[0]
                tr_pr = hdr._tr.get_or_add_trPr()
                th = OxmlElement("w:tblHeader")
                th.set(qn("w:val"), "true")
                tr_pr.append(th)
                for j, h in enumerate(T.headers[lang]):
                    add_marked(hdr.cells[j].paragraphs[0], T.mark(h), cell_pt, base_bold=True)
                for r in T.rows:
                    row = tbl.add_row()
                    if isinstance(r, tuple):
                        merged = row.cells[0].merge(row.cells[ncols - 1])
                        add_marked(merged.paragraphs[0], bt.unbold(r[1].text[lang]), cell_pt, base_bold=True)
                        continue
                    for j, c in enumerate(r):
                        add_marked(row.cells[j].paragraphs[0], T.mark(c.md(lang)), cell_pt)
                for k, (_key, txt) in enumerate(T.notes):
                    np_ = doc.add_paragraph()
                    style_run(np_.add_run("abcdefghijklmnopqrstuvwxyz"[k]), body_pt, sup_=True)
                    add_marked(np_, " " + txt[lang], body_pt)
        ensure_orientation(False)
        heading(SOURCES[lang][0], body_pt + 2)
        for para in SOURCES[lang][1:]:
            add_marked(doc.add_paragraph(), para, body_pt)
        doc.save(str(path))


def build(lang: str, path: Path) -> None:
    N = common.load_numbers()
    P = bt.load_prespec()
    write_combined(lang, path, bt.build_all(N, P))


def make_zip(path: Path = ZIP) -> list[str]:
    """配布用の zip。out/ の図と凡例文、表の md・csv・docx、まとめの docx、README を入れる。"""
    names = []
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for p in sorted(OUT.glob("*")):
            if p.suffix in (".pdf", ".svg", ".png", ".txt"):
                z.write(p, f"figtab/out/{p.name}")
                names.append(p.name)
        for p in [HERE / "tables_ja.md", HERE / "tables_en.md", HERE / "tables_ja.docx", HERE / "tables_en.docx",
                  DOCX["ja"], DOCX["en"], HERE / "README.md"]:
            z.write(p, f"figtab/{p.name}")
            names.append(p.name)
        for p in sorted((HERE / "csv").glob("*.csv")):
            z.write(p, f"figtab/csv/{p.name}")
            names.append(p.name)
    return names


def docx_summary(path: Path) -> dict:
    from docx import Document
    d = Document(str(path))
    texts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            for c in row.cells:
                texts.append(c.text)
    return {"n_pictures": len(d.inline_shapes), "n_tables": len(d.tables), "text": "\n".join(texts)}


def selftest() -> int:
    ok_all = True

    def rep(name, ok, detail=""):
        nonlocal ok_all
        ok_all &= bool(ok)
        print(f"  {'PASS' if ok else 'FAIL'}  {name}" + (f"  ({detail})" if detail and not ok else ""))

    banned, pragma = bt._banned_terms()
    n_fig = sum(len(p["figures"]) for p in PARTS)
    n_tab = sum(len(p["tables"]) for p in PARTS)
    with tempfile.TemporaryDirectory() as td:
        for lang in ("ja", "en"):
            print(f"[{lang}]")
            path = Path(td) / f"t_{lang}.docx"
            build(lang, path)
            s = docx_summary(path)
            rep(f"図が {n_fig} 点入っている", s["n_pictures"] == n_fig, f"{s['n_pictures']}")
            rep(f"表が {n_tab} 点入っている", s["n_tables"] == n_tab, f"{s['n_tables']}")
            tabs = bt.build_all(common.load_numbers(), bt.load_prespec())
            rep("build_tables の表がすべて 3 部のどれかに割り当てられている",
                sorted(T.tid for T in tabs) == sorted(t for p in PARTS for t in p["tables"]))
            missing = [name for p in PARTS for name in p["figures"]
                       if legend_paragraphs(name, lang)[0] not in s["text"]]
            rep("各図の凡例文の 1 行目が文書にある", not missing, f"{missing}")
            rep("段の注釈が文書にある", common.STAGE_NOTE[lang] in s["text"])
            if banned is None:
                print("  （用語検査器が無いので禁止語の確認は飛ばした）")
            else:
                hits = bt._banned_hits(banned, pragma, s["text"])
                rep("文書に禁止語が無い（検査器の表）", not hits, f"{hits[:3]}")
            for name in [n for p in PARTS for n in p["figures"]]:
                rep_ok = (OUT / f"{name}_{lang}.png").exists() and (OUT / f"{name}_legend_{lang}.txt").exists()
                if not rep_ok:
                    rep(f"{name} の PNG と凡例文がある", False, name)
            widths = {n: common.png_size_mm(OUT / f"{n}_{lang}.png")[0] for p in PARTS for n in p["figures"]}
            # 許容は 0.1 mm（600 dpi で約 2 画素）。bbox="tight" で書き出すときの画素の丸めの分
            wide = {n: round(w, 2) for n, w in widths.items() if w > common.FIG_MAX_WIDTH_MM + 0.1}
            rep(f"図の PNG（600 dpi）の幅がすべて {common.FIG_MAX_WIDTH_MM:g} mm 以下", not wide, f"{wide}")
            from docx import Document
            shapes = Document(str(path)).inline_shapes
            over = [round(sh.width / 36000, 1) for sh in shapes if sh.width / 36000 > FIG_WIDTH_MM + 0.1]
            rep(f"文書に貼った図の幅がすべて {FIG_WIDTH_MM:g} mm 以下", not over, f"{over}")
            intro = "".join(INTRO[lang]).replace("{n_fig}", str(n_fig)).replace("{n_tab}", str(n_tab))
            rep("前書きの図と表の数が 3 部の構成と一致", (f"図 {n_fig} 点と表 {n_tab} 点" in intro) if lang == "ja"
                else (f"{n_fig} figures and {n_tab} tables" in intro))
    print("RESULT", "PASS" if ok_all else "FAIL")
    return 0 if ok_all else 1


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--selftest", action="store_true")
    ap.add_argument("--zip", action="store_true", help="配布用の zip も作る")
    a = ap.parse_args()
    if a.selftest:
        return selftest()
    for lang, path in DOCX.items():
        build(lang, path)
        s = docx_summary(path)
        print(f"{path}  図 {s['n_pictures']}・表 {s['n_tables']}  {path.stat().st_size / 1e6:.1f} MB")
    if a.zip:
        names = make_zip()
        print(f"{ZIP}  {len(names)} ファイル  {ZIP.stat().st_size / 1e6:.1f} MB")
    return 0


if __name__ == "__main__":
    sys.exit(main())
