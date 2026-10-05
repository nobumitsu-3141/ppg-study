#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""論文2 表集 ― 和文・英文 Markdown、CSV、docx を同じ JSON から出す。

表の並び（この図表集の番号。括弧は `../02_tables.md` の番号）
    表1   事前に決めた解析項目の一覧（`data/prespec_chronology.json` prespecified_items）
    表2   判定表（02 表1）
    表3   因子ごとの主効果（02 表2）、表3b 1 因子掃引（02 表2b）、表3c 真の伝播時間（02 表2c）
    表4   基底関数と成分数（02 表3。探索・事後）
    表5a  公表された条件の再現（02 表4）、表5b 逸脱表（02 表4 逸脱表）
    表6a  波形の型で分けた関連（02 表5）、表6b 真値の層内の幅（02 表5b）
    表7   診断と改善案の対応（`prespec_chronology.json` variants_table6 ＋ 02 表6 の数値）
    表8   拡張期の下降の扱いを変えたときの型3 の関連（02 表6）
    表9   B 段と雑音（02 表6c）
    補足表 S1  端に達した母数（02 表6b）
    補足表 S2a 線形分離の後進波との照合（02 表7a）、S2b 圧波形に当てた結果（02 表7b）

数値は `data/paper2_numbers.json`（`data/extract_tables.py` が `../02_tables.md` を機械で読んだもの）からだけ取り、
手で打たない。文章の列（表1 の規則・表7 の変更点など）は `data/prespec_chronology.json` から取り、
その中の数値は自己検査で出典の文字列と照合する。名称（和・英）は `data/labels.json` をそのまま使う。
判定（成立・不成立）は `roadmap_v1.md` §9 のまま。表4 以降は探索・事後であり、脚注にその旨を書く。
段（A・B・C）を出す表には毎回 1 行の注釈を置く（CLAUDE.md §3）。

使い方
    python3 build_tables.py                 和文・英文の両方を書く（tables_ja.md・tables_en.md・csv/・docx）
    python3 build_tables.py --lang en       英文だけ
    python3 build_tables.py --selftest      出力を読み直して JSON と照合する（数値・太字・禁止語・docx・csv）
"""
from __future__ import annotations

import argparse
import csv
import importlib.util
import json
import re
import subprocess
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
from common import load_numbers          # noqa: E402

REPO = HERE.parents[3]
NUMBERS = HERE / "data" / "paper2_numbers.json"
PRESPEC = HERE / "data" / "prespec_chronology.json"
LABELS = json.loads((HERE / "data" / "labels.json").read_text(encoding="utf-8"))
OUT_LEGEND = HERE / "out"

LANGS = ("ja", "en")
_LEGACY_EN_TERM = "landmark"   # 用語の引用（英文の地の文にこの語が無いことを自己検査で確かめるための文字列）
MAX_JA_CHARS_T7 = 60      # 表7 の文章の列の上限（和文の文字数）
MAX_EN_WORDS_T7 = 20      # 同（英文の語数）

# ---------------------------------------------------------------------------
# 文字と数値の小道具
# ---------------------------------------------------------------------------

_SUPS = str.maketrans("⁰¹²³⁴⁵⁶⁷⁸⁹", "0123456789")
_NUM_RE = re.compile(r"(?<![A-Za-z\d])[+\-−]?\d[\d,]*(?:\.\d+)?(?:e[+-]?\d+)?")
# 母数名・記号に付いた数字（α1・σ2・R1_d・p1・Am_p1・T2・P6 など）は数値として数えない
_IDENT_RE = re.compile(r"[A-Za-z_Α-ω]+\d+[A-Za-z_]*")
_SUP_TAG_RE = re.compile(r"<sup>[a-z](?:, [a-z])*</sup>")


def nums_in(text: str) -> list[float]:
    """文字列に現れる数値を順に返す（自己検査用。`**`・脚注の印・母数名の添字は除く）。"""
    t = _SUP_TAG_RE.sub("", text).replace("**", "")
    t = _IDENT_RE.sub(" ", t)
    t = re.sub(r"×10([⁰¹²³⁴⁵⁶⁷⁸⁹]+)", lambda m: "e" + m.group(1).translate(_SUPS), t)
    out = []
    for m in _NUM_RE.finditer(t):
        s = m.group(0).replace(",", "").replace("−", "-")
        try:
            out.append(float(s))
        except ValueError:
            pass
    return out


def same_nums(a: list[float], b: list[float], tol: float = 1e-9) -> bool:
    return len(a) == len(b) and all(abs(x - y) <= tol for x, y in zip(a, b))


def unbold(s: str) -> str:
    return s.replace("**", "")


def plain(s: str) -> str:
    """Markdown の印（太字・脚注の上付き）を外した文字列（CSV 用）。"""
    return _SUP_TAG_RE.sub("", unbold(s)).replace("\\|", "|")


_FW_EN = str.maketrans({"（": "(", "）": ")", "／": "/", "〜": "–", "　": " "})


def en_from_raw(raw: str) -> str:
    """02_tables.md の和文セル（raw）を英文の表記に直す。太字の印はそのまま残す。"""
    s = raw.translate(_FW_EN)
    s = s.replace("(", " (")
    s = re.sub(r"\s+\(", " (", s)
    rules = [
        ("A・B は判定不能", "tiers A and B not evaluable"),
        ("採択 ", "accepted "),
        ("、C ", "; tier C "),
        ("。", "; "),
        ("; A ", "; tier A "),
        (" 層)", " strata)"),
        ("未検証", "not evaluated"),
        ("(基準)", "(reference)"),
        ("(同左)", "(as at left)"),
    ]
    for a, b in rules:
        s = s.replace(a, b)
    s = re.sub(r"^(\*\*)?A (\d)", r"\1tier A \2", s)
    s = re.sub(r"\s+", " ", s).strip()
    s = s.replace("( ", "(").replace(" )", ")")
    return s


def md_escape(s: str) -> str:
    return s.replace("|", "\\|")


def sup(letter: str) -> str:
    return f"<sup>{letter}</sup>"


def fmt_int(n) -> str:
    return f"{int(n):,}"


def find_num(text: str, pattern: str, group: int = 1) -> str:
    """文章（JSON の preamble・postscript）から数値の文字列を取り出す。無ければ止まる。"""
    m = re.search(pattern, text)
    if not m:
        raise ValueError(f"文章から数値を取り出せない: {pattern!r} in {text[:60]!r}")
    return m.group(group)


# ---------------------------------------------------------------------------
# セルと表の型
# ---------------------------------------------------------------------------

class Cell:
    """1 つのセル。`text` は言語ごとの Markdown 文字列（`**` で太字）。`nums` は文字列に現れるべき数値
    （JSON の値。None なら数値の照合をしない）。`src` は文章の列の出典文字列（数値の部分集合の照合に使う）。"""

    __slots__ = ("text", "nums", "bold", "src", "kind", "csv")

    def __init__(self, text: dict, nums=None, bold: bool = False, src: str | None = None, kind: str = "data",
                 csv_text: dict | None = None):
        self.text = text
        self.nums = nums
        self.bold = bold
        self.src = src
        self.kind = kind
        self.csv = csv_text          # CSV に書く文字列（None なら text の印を外したもの）

    def md(self, lang: str) -> str:
        s = self.text[lang]
        if self.bold and not s.startswith("**"):
            s = f"**{s}**"
        return md_escape(s) if "\\|" not in s else s


def label_cell(ja: str, en: str, bold: bool = False, nums=None) -> Cell:
    return Cell({"ja": ja, "en": en}, nums=nums, bold=bold, kind="label")


def text_cell(ja: str, en: str, src: str | None = None, bold: bool = False) -> Cell:
    return Cell({"ja": ja, "en": en}, nums=None, bold=bold, src=src, kind="text")


class Table:
    def __init__(self, tid: str, number: dict, caption: dict, headers: dict, notes: list,
                 landscape: bool = False, small: bool = False, csv_name: str | None = None):
        self.tid = tid                   # 機械用の名（csv のファイル名）
        self.number = number             # {"ja": "表1", "en": "Table 1"}
        self.caption = caption           # 言語ごとの見出し文（脚注の印は {a} のように書く）
        self.headers = headers           # 言語ごとの見出しの列（印を含めてよい）
        self.notes = notes               # [(key, {"ja":..., "en":...}), ...]（a, b, c … の順）
        self.rows: list = []             # 行は Cell の列、または ("group", Cell)
        self.landscape = landscape
        self.small = small
        self.csv_name = csv_name or tid

    def add(self, *cells: Cell):
        self.rows.append(list(cells))

    def group(self, cell: Cell):
        self.rows.append(("group", cell))

    def data_rows(self):
        return [r for r in self.rows if not (isinstance(r, tuple) and r[0] == "group")]

    def letter(self, key: str) -> str:
        for i, (k, _t) in enumerate(self.notes):
            if k == key:
                return "abcdefghijklmnopqrstuvwxyz"[i]
        raise KeyError(key)

    def mark(self, s: str) -> str:
        """文字列の中の {key} を脚注の上付き文字に置き換える。"""
        s = re.sub(r"\{([a-z_0-9]+)\}", lambda m: sup(self.letter(m.group(1))), s)
        return s.replace("</sup><sup>", ", ")

    def ncols(self, lang: str) -> int:
        return len(self.headers[lang])


# 段の注釈と探索・事後の注記（labels.json。図とも共有）
STAGE_NOTE = LABELS["stage_note"]
POSTHOC_NOTE = LABELS["posthoc_note"]
NOT_FOR_DECISION = {"ja": "判定には用いない。", "en": "Not used for the decision."}
STAGE_NAME = {k: {"ja": v["ja"], "en": v["en"]} for k, v in LABELS["stages"].items()}


def rho_cell(c: dict, lang_text: dict | None = None) -> Cell:
    """`kind: rho` のセル（0.223（6/6）、—、未検証、A/C の段付きなど）。"""
    raw = c["raw"]
    if c.get("dash"):
        return Cell({"ja": "—", "en": "—"}, nums=[], bold=False)
    if c.get("note") and c.get("rho") is None:
        return Cell({"ja": raw, "en": en_from_raw(raw)}, nums=[], bold=c.get("bold", False))
    nums: list[float] = []
    if c.get("rho") is not None:
        nums.append(c["rho"])
        if c.get("strata_pass") is not None:
            nums += [c["strata_pass"], c["strata_total"]]
    if "stages" in c:
        nums = []
        if c.get("adopted"):
            nums.append(c["adopted"]["n"])
            if c["adopted"].get("of") is not None:
                nums.append(c["adopted"]["of"])
        for st in ("A", "B", "C"):
            sc = c["stages"].get(st)
            if sc and sc.get("rho") is not None:
                nums.append(sc["rho"])
                if sc.get("strata_pass") is not None:
                    nums += [sc["strata_pass"], sc["strata_total"]]
    text = {"ja": raw, "en": en_from_raw(raw)}
    if lang_text:
        text.update(lang_text)
    return Cell(text, nums=nums, bold=c.get("bold", False))


def value_cell(c: dict, suffix: str = "") -> Cell:
    """`kind: num / pct / int / ms` のセル。文字列は raw をそのまま使う。"""
    raw = c["raw"]
    if c.get("dash") or c.get("value") is None:
        return Cell({"ja": raw, "en": raw.translate(_FW_EN)}, nums=[], bold=c.get("bold", False))
    return Cell({"ja": raw, "en": en_from_raw(raw)}, nums=[c["value"]], bold=c.get("bold", False))


def posthoc_notes(table_key_ja: str) -> list:
    """表4 以降に共通の脚注（探索・事後、判定には用いない）。"""
    # labels.json の注記（探索・事後。判定には用いない。主解析に用いるには新たな事前登録が要る）をそのまま使う
    return [("posthoc", {"ja": POSTHOC_NOTE["ja"], "en": POSTHOC_NOTE["en"]})]


def entries_in(src: str) -> list[str]:
    """「追記144・146・147」「追記16〜18」のような列挙から追記番号を順に取り出す。"""
    out: list[str] = []
    for m in re.finditer(r"追記([\d・〜]+)", src):
        for part in m.group(1).strip("・〜").split("・"):
            if "〜" in part:
                a, b = part.split("〜")
                out += [str(i) for i in range(int(a), int(b) + 1)]
            elif part:
                out.append(part)
    seen: list[str] = []
    for e in out:
        if e not in seen:
            seen.append(e)
    return seen


def table_name(t: dict) -> dict:
    name = t["name"]
    if name == "逸脱表":
        return {"ja": "表4 の逸脱表", "en": "the deviation table under table 4"}
    return {"ja": name, "en": name.replace("表", "table ")}


def source_note(t: dict, extra_ja: str = "", extra_en: str = "") -> dict:
    """JSON の出典から脚注を組む（02_tables の表番号・台本・結果ファイル・追記番号）。
    文書の冒頭にしか出典が無い表は、その表に当たる部分（source_segments）だけを使う。"""
    if t.get("source_kind") == "文書の冒頭（出どころ）":
        src = " ".join(t.get("source_segments") or [])
    else:
        src = t.get("source") or t.get("source_paragraph") or ""
    scripts = sorted(set(re.findall(r"`analysis/scripts/(\d+)_[^`]+`", src)))
    files = re.findall(r"`([^`]+\.txt)`", src)
    entries = entries_in(src)
    nm = table_name(t)
    ja = f"出典: `02_tables.md` {nm['ja']}"
    en = f"Source: {nm['en']} of `02_tables.md`"
    if scripts:
        ja += "。台本 " + "・".join(f"{x}番" for x in scripts)
        en += "; script" + ("s " if len(scripts) > 1 else " ") + ", ".join(scripts)
    if files:
        ja += "（" + "・".join(f"`{f}`" for f in files) + "）"
        en += " (" + ", ".join(f"`{f}`" for f in files) + ")"
    if entries:
        ja += "。lab_log 追記" + "・".join(entries)
        en += "; lab_log entr" + ("ies " if len(entries) > 1 else "y ") + ", ".join(entries)
    ja += "。" + extra_ja
    en += ". " + extra_en
    return {"ja": ja.strip(), "en": en.strip()}


# ---------------------------------------------------------------------------
# 表1  事前に決めた解析項目の一覧（prespec_chronology.json ＋ 02_tables の数値）
# ---------------------------------------------------------------------------

VERDICT_EN = {"不成立": "fail", "成立": "pass", "合格": "pass"}

# 行の順（課題の指定）。鍵は prespecified_items の name_ja に含まれる語
T1_ORDER = ["凍結版", "第2版 歪みガウス", "第2版 ガンマ", "特徴点法（Charlton", "早期振幅比", "陽性対照",
            "p1 基準", "DPS", "波形の型の分布", "27番", "表2 振った", "表2b", "表2c"]

T1_EN = {
    "凍結版": dict(
        name="Frozen decomposition (two skew-Gaussian components): ΔT × aortic PWV (Q1), RI × peripheral vascular resistance (Q2)",
        role="Primary questions Q1 and Q2 (table 2, row 1)",
        rule="Within-stratum (six strata) Spearman ρ with the predicted sign in every stratum (ΔT negative, RI positive) "
             "and median |ρ| ≥ 0.3; strata of at least 8 subjects; pooled correlations are reference values only",
        frozen="2026-09-03 (before script 20 was run); the question itself on 2026-08-30 (roadmap §3, study 0)",
        verdict="fail",
        source="lab_log 2026-09-03 (prespecified criteria of study 0; results of study 0); roadmap §8, §9; reproduced in entry 12"),
    "第2版 歪みガウス": dict(
        name="Rebuilt decomposition, skew-Gaussian route (primary route; src/pda2.py)",
        role="Primary question: was the failure of the frozen version an implementation problem? (table 2, row 2)",
        rule="Definition of 'pass' 1–3: both questions pass in all three tiers A, B and C, without an asterisk; "
             "if script 27 splits the verdict, 'pass' is not written",
        frozen="2026-09-04 (gate0_rules_v2, pda2_thresholds_v2); pda2 version 048d2b43bb05 fixed on 2026-09-05",
        verdict="fail (tiers A and B not evaluable; tier C fails). Decomposition is not adopted as a vascular index",
        source="gate0_rules_v2.md; pda2_final_summary.md §1–2; entry 12 §1–3; roadmap §9"),
    "第2版 ガンマ": dict(
        name="Rebuilt decomposition, gamma route (reported alongside)",
        role="Robustness check (table 2, row 3)",
        rule="'A pass of the gamma route alone is not written as a pass of PDA'; the proportion of parameters on a bound "
             "and the script-27 part-B (refitting) 'widen the search bounds' verdict are reported alongside",
        frozen="2026-09-04",
        verdict="fail (the verdict splits across thresholds and conditions in script 27)",
        source="gate0_rules_v2.md; entries 12, 13"),
    "特徴点法（Charlton": dict(
        name="Fiducial-point analysis (indices supplied with PWDB, same waveforms): ΔT, RI; secondary SI, AGI_mod, AI, ΔT × carotid–femoral PWV",
        role="Control: only the index construction differs (table 2, row 4)",
        rule="Same criterion imported from script 20. Reading fixed before the run: 'fiducial-point analysis also fails → the failure "
             "is not specific to PDA'; 'only fiducial-point analysis passes → the failure is specific to PDA'; "
             "'model-derived PTT not strongly negative → doubt the comparison itself'",
        frozen="2026-09-03 (before script 23 was run)",
        verdict="pass (AI alone fails). The failure is specific to decomposition; the full withdrawal was retracted",
        source="lab_log 2026-09-03 (control experiment; correction of the verdict); roadmap §8; reproduced in entry 12"),
    "早期振幅比": dict(
        name="Early amplitude ratio Am_b/Am_p1 (Hellqvist 2024) × aortic PWV",
        role="Fourth method, exploratory (table 2, row 5)",
        rule="'If it passes and is at least equal to the median of fiducial-point ΔT (a difference within 0.05 counts as equal), "
             "the candidate primary index of study 2 switches to Am_b/Am_p1 and ΔT becomes secondary'; "
             "'a 5–95% width below 0.05 is read as no discrimination'",
        frozen="2026-09-04",
        verdict="pass (fails against carotid–femoral PWV). Candidate primary index of study 2 switched to Am_b/Am_p1",
        source="gate0_rules_v2.md; entry 12 §3-4"),
    "陽性対照": dict(
        name="Positive control: model-derived pulse transit time × aortic PWV",
        role="Check of the analysis pipeline (table 2, row 6)",
        rule="2026-09-03: if not strongly negative, doubt the comparison itself. 2026-09-04: the whole table is void unless "
             "median |ρ| ≥ 0.50 with every stratum negative",
        frozen="2026-09-03 (script 23); 2026-09-04 (quantified in gate0_rules_v2)",
        verdict="pass; the table is valid",
        source="lab_log 2026-09-03 (correction of the verdict); gate0_rules_v2.md; entries 11, 12 §0"),
    "p1 基準": dict(
        name="ΔT with p1 (Hellqvist's systolic peak) as the systolic reference",
        role="Exploratory (note to table 2)",
        rule="'If p1-based ΔT and our own fiducial-point ΔT differ greatly in type 3, replace the systolic reference of "
             "fiducial-point ΔT by p1'",
        frozen="2026-09-04",
        verdict="fail; p1 is not adopted (the rule's row does not apply)",
        source="gate0_rules_v2.md; entry 12 §3-5"),
    "DPS": dict(
        name="DPS (difference of component widths, Goswami 2010)",
        role="Descriptive",
        rule="'If the DPS of the rebuilt skew-Gaussian route is monotonic in the true value, add it as a secondary index of "
             "study 2 (sign not decided post hoc; reported as description)'",
        frozen="2026-09-04",
        verdict="descriptive only (the rule presupposes an accepted decomposition, so it is not applied)",
        source="gate0_rules_v2.md; entry 12 §3-7"),
    "波形の型の分布": dict(
        name="Waveform-type distribution (script 26, table 2) and handling of types 3 and 4",
        role="Condition for reading the verdict",
        rule="'If type 3 is frequent, the PDA verdict rests on type-1 beats and type 3 is read with fiducial points, p1 and the "
             "early amplitude ratio; report n of type 3 and the script-27 part-A \"include type 3\" verdict'; "
             "'no ΔT-type index can be defined for type 4; report its n and discuss it with the early amplitude ratio only'",
        frozen="2026-09-04",
        verdict="read as the rule prescribes: the decomposition verdict rests on type-1 beats; type 3 read with fiducial points, "
                "p1 and the early amplitude ratio; type 4 with the early amplitude ratio only",
        source="gate0_rules_v2.md; entry 12 §1, §3-6"),
    "27番": dict(
        # 27番の A 層・B 層（感度解析の 2 つの部分）は被験者の段（tier）ではないので part と訳す
        name="Threshold sensitivity analysis (script 27): part A, acceptance recomputed without refitting; part B, refitted on a fixed "
             "subset of 624 subjects under 14 conditions",
        role="Check that the conclusion is not a product of the thresholds",
        rule="'If every threshold gives the same answer, the conclusion is not a product of the thresholds; if it splits, write the "
             "range and do not write pass'. The part-B subset is fixed to subj_no % 7 == 0 and is not re-chosen after seeing results",
        frozen="2026-09-04 (lab_log: thresholds frozen before the decision test)",
        verdict="reflected in the two rebuilt-version rows of table 2",
        source="pda2_thresholds_v2.md; entry 12 §2"),
    "表2 振った": dict(
        name="Table 3: within-stratum main effect of each varied factor (descriptive)",
        role="Mechanism, descriptive",
        rule="'If the main factor moving ΔT is PWV and the main factor moving RI is MBP, the concept holds'",
        frozen="2026-09-03",
        verdict="not as conceived: for decomposition ΔT the heart-rate effect is of the same order as that of pulse wave velocity, "
                "and the main factor of RI is not mean arterial pressure (MBP)",
        source="lab_log 2026-09-03 (results of study 0; correction of the verdict); entry 12; table 3 (02_tables table 2)"),
    "表2b": dict(
        name="Table 3b: one-factor-at-a-time sweep (descriptive)",
        role="Mechanism, descriptive",
        rule="As above (values of the subjects in which one factor alone is varied)",
        frozen="2026-09-03",
        verdict="the response is not monotonic (the correspondence of the components exchanges)",
        source="lab_log 2026-09-03 (results of study 0); table 3b (02_tables table 2b)"),
    "表2c": dict(
        name="Table 3c: true transit times from the model onset-time table (descriptive)",
        role="True values, descriptive",
        rule="'Compare the width of the true radial→finger transit time with the within-case SD of T2−T1 in VitalDB (≈18 ms)'",
        frozen="2026-09-03",
        verdict="the width of the true radial→finger transit time is narrower than the VitalDB within-case SD; decomposition ΔT "
                "barely tracks the true transit time",
        source="lab_log 2026-09-03 (results of study 0); table 3c (02_tables table 2c)"),
}

T1_ROLE_JA = {
    "凍結版": "主要な問い Q1・Q2（表2 行1）",
    "第2版 歪みガウス": "主要な問い（凍結版の失敗が実装の問題かを問う。表2 行2）",
    "第2版 ガンマ": "頑健性の確認（表2 行3）",
    "特徴点法（Charlton": "対照（指標の作り方だけを変える。表2 行4）",
    "早期振幅比": "4 つめの手法・探索（表2 行5）",
    "陽性対照": "処理系の検査（表2 行6）",
    "p1 基準": "探索（表2 の注）",
    "DPS": "記述",
    "波形の型の分布": "判定の読み方の条件",
    "27番": "結論が閾値の産物でないかの検査",
    "表2 振った": "機構の記述",
    "表2b": "機構の記述",
    "表2c": "真値の記述",
}

T1_VERDICT_JA = {
    "凍結版": "不成立",
    "第2版 歪みガウス": "不成立（A・B 判定不能・C 不成立）。分解法を血管指標として用いる方針は取らない",
    "第2版 ガンマ": "不成立（27番で閾値・条件により判定が割れる）",
    "特徴点法（Charlton": "成立（AI のみ不成立）。失敗は分解に固有と訂正し、全面撤退を取り消した",
    "早期振幅比": "成立（頸大腿PWV に対しては不成立）。研究2 の主指標候補を Am_b/Am_p1 に切り替えた",
    "陽性対照": "合格。表は有効",
    "p1 基準": "不成立。p1 は採用しない（該当行に該当せず）",
    "DPS": "記述のみ（採用された分解を前提とする行なので適用しない）",
    "波形の型の分布": "該当行どおりに読んだ: 分解法の判定は型1 の拍で下し、型3 は特徴点法・p1・早期振幅比で、型4 は早期振幅比だけで読む",
    "27番": "表2 の第2版の 2 行に反映",
    "表2 振った": "概念どおりではない: 分解法 ΔT は心拍数の効きが脈波伝播速度と同じ桁で、RI の主因子は平均血圧ではない",
    "表2b": "応答が単調でない（成分の対応が入れ替わる）",
    "表2c": "橈骨→指尖の真の伝播時間の幅は VitalDB の症例内 SD より狭い。分解法 ΔT は真の伝播時間をほとんど追わない",
}


def t1_results(N: dict) -> dict:
    """表1 の「結果」列を 02_tables の数値（JSON）から組む。鍵は T1_ORDER の語。"""
    t = N["tables"]
    r1 = {r["id"]: r for r in t["表1"]["rows"]}
    post = "\n".join(t["表1"]["postscript"])
    e = en_from_raw

    def rc(c):
        return rho_cell(c)

    out = {}
    c = r1["pda_frozen"]
    out["凍結版"] = Cell({"ja": f"ΔT {c['dt_pwv']['raw']}・RI {c['ri_pvr']['raw']}",
                         "en": f"ΔT {e(c['dt_pwv']['raw'])}; RI {e(c['ri_pvr']['raw'])}"},
                        nums=rc(c["dt_pwv"]).nums + rc(c["ri_pvr"]).nums)
    c = r1["pda_v2_skewgauss"]
    out["第2版 歪みガウス"] = Cell({"ja": f"ΔT: {c['dt_pwv']['raw']}。RI: {c['ri_pvr']['raw']}",
                                  "en": f"ΔT: {e(c['dt_pwv']['raw'])}. RI: {e(c['ri_pvr']['raw'])}"},
                                 nums=rc(c["dt_pwv"]).nums + rc(c["ri_pvr"]).nums)
    c = r1["pda_v2_gamma"]
    out["第2版 ガンマ"] = Cell({"ja": f"ΔT: {c['dt_pwv']['raw']}。RI: {c['ri_pvr']['raw']}",
                              "en": f"ΔT: {e(c['dt_pwv']['raw'])}. RI: {e(c['ri_pvr']['raw'])}"},
                             nums=rc(c["dt_pwv"]).nums + rc(c["ri_pvr"]).nums)
    c = r1["landmark"]
    si = find_num(post, r"スティフネス指標 ([\d.]+)・2次微分の加齢指数 ([\d.]+) が成立し、\s*増大係数 ([\d.]+) は不成立", 1)
    agi = find_num(post, r"スティフネス指標 ([\d.]+)・2次微分の加齢指数 ([\d.]+) が成立し", 2)
    ai = find_num(post, r"増大係数 ([\d.]+) は不成立", 1)
    out["特徴点法（Charlton"] = Cell(
        {"ja": f"ΔT {unbold(c['dt_pwv']['raw'])}・RI {unbold(c['ri_pvr']['raw'])}。副次: SI {si}・AGI_mod {agi} は成立、AI {ai} は不成立",
         "en": f"ΔT {e(unbold(c['dt_pwv']['raw']))}; RI {e(unbold(c['ri_pvr']['raw']))}. Secondary: SI {si} and AGI_mod {agi} pass, AI {ai} fails"},
        nums=rc(c["dt_pwv"]).nums + rc(c["ri_pvr"]).nums + [float(si), float(agi), float(ai)])
    c = r1["amp_ratio"]
    m = re.search(r"頸大腿脈波伝播速度に対しては ([\d.]+)（(\d)/(\d)）で成立しない", post)
    if not m:
        raise ValueError("表1 の注から頸大腿PWV の値を取り出せない")
    out["早期振幅比"] = Cell(
        {"ja": f"{unbold(c['dt_pwv']['raw'])}。頸大腿PWV に対しては {m.group(1)}（{m.group(2)}/{m.group(3)}）で不成立",
         "en": f"{e(unbold(c['dt_pwv']['raw']))}; against carotid–femoral PWV {m.group(1)} ({m.group(2)}/{m.group(3)}), a fail"},
        nums=rc(c["dt_pwv"]).nums + [float(m.group(1)), float(m.group(2)), float(m.group(3))])
    c = r1["ptt_control"]
    out["陽性対照"] = Cell({"ja": c["dt_pwv"]["raw"], "en": e(c["dt_pwv"]["raw"])}, nums=rc(c["dt_pwv"]).nums)
    m = re.search(r"p1 に置き換えた ΔT は ([\d.]+)（(\d)/(\d)）で、\s*同一の ([\d,]+) 名における自前の特徴点由来 ΔT ([\d.]+) に劣った", post)
    if not m:
        raise ValueError("表1 の注から p1 基準 ΔT の値を取り出せない")
    out["p1 基準"] = Cell(
        {"ja": f"{m.group(1)}（{m.group(2)}/{m.group(3)}）。同一の {m.group(4)} 名における自前の特徴点由来 ΔT {m.group(5)} に劣る",
         "en": f"{m.group(1)} ({m.group(2)}/{m.group(3)}); in the same {m.group(4)} subjects, below our own fiducial-point ΔT {m.group(5)}"},
        nums=[float(m.group(1)), float(m.group(2)), float(m.group(3)), float(m.group(4).replace(",", "")), float(m.group(5))])
    n_acc = r1["pda_v2_skewgauss"]["dt_pwv"]["adopted"]["n"]
    out["DPS"] = Cell({"ja": f"主たる経路（歪みガウス）の採択は {n_acc} 拍で、A 段は判定不能。C 段の値は記述にとどめる",
                       "en": f"The primary (skew-Gaussian) route accepted {n_acc} beats, so tier A is not evaluable; tier-C values are reported as description only"},
                      nums=[float(n_acc)])
    t5 = {r["id"]: r for r in t["表5"]["rows"]}
    n_all = N["meta"]["n_subjects_decision_test"]
    parts_ja, parts_en, nums = [], [], []
    for tid, digit in (("type1", 1), ("type3", 3), ("type4", 4)):
        n = t5[tid]["n"]
        pct = round(100.0 * n / n_all, 1)
        parts_ja.append(f"型{digit} {fmt_int(n)}名（{pct:.1f}%）")
        parts_en.append(f"type {digit} {fmt_int(n)} ({pct:.1f}%)")
        nums += [float(digit), float(n), pct]
    out["波形の型の分布"] = Cell({"ja": "・".join(parts_ja), "en": "; ".join(parts_en)}, nums=nums)
    # 27番の「A 層・B 層」は感度解析の 2 つの部分（当てはめ直し無し／有り）で、被験者の段（tier A・B・C）ではない。
    out["27番"] = Cell({"ja": "歪みガウス経路はどの閾値でも不成立。ガンマ経路は A 層・B 層とも閾値・条件で判定が割れる",
                       "en": "The skew-Gaussian route fails at every threshold; the gamma route splits in both part A and part B"},
                      nums=[])
    r2 = {r["id"]: r for r in t["表2"]["rows"]}
    pd_ = lambda f: unbold(r2[f]["pda_dt"]["raw"])      # noqa: E731
    ld_ = lambda f: unbold(r2[f]["landmark_dt"]["raw"])  # noqa: E731
    pr_ = lambda f: unbold(r2[f]["pda_ri"]["raw"])      # noqa: E731
    nums = []
    for f in ("pwv", "heart_rate", "aortic_diameter"):
        nums += [r2[f]["pda_dt"]["effect_pct"], r2[f]["pda_dt"]["rho"]]
    for f in ("pwv", "heart_rate"):
        nums += [r2[f]["landmark_dt"]["effect_pct"]]
    for f in ("heart_rate", "pwv", "map"):
        nums += [r2[f]["pda_ri"]["effect_pct"], r2[f]["pda_ri"]["rho"]]
    out["表2 振った"] = Cell(
        {"ja": f"分解法 ΔT: 脈波伝播速度 {pd_('pwv')}・心拍数 {pd_('heart_rate')}・大動脈径 {pd_('aortic_diameter')}％。"
               f"特徴点法 ΔT: 脈波伝播速度 {ld_('pwv')}・心拍数 {ld_('heart_rate')}％。"
               f"分解法 RI: 心拍数 {pr_('heart_rate')}・脈波伝播速度 {pr_('pwv')}・平均血圧 {pr_('map')}％",
         "en": f"Decomposition ΔT: PWV {e(pd_('pwv'))}, heart rate {e(pd_('heart_rate'))}, aortic diameter {e(pd_('aortic_diameter'))}%. "
               f"Fiducial-point ΔT: PWV {ld_('pwv')}, heart rate {ld_('heart_rate')}%. "
               f"Decomposition RI: heart rate {e(pr_('heart_rate'))}, PWV {e(pr_('pwv'))}, mean arterial pressure {e(pr_('map'))}%"},
        nums=nums)
    r2b = {r["id"]: r for r in t["表2b"]["rows"]}["pwv"]
    keys = ("dt_minus1sd", "dt_base", "dt_plus1sd", "ri_minus1sd", "ri_base", "ri_plus1sd")
    v = [unbold(r2b[k]["raw"]) for k in keys]
    # 02 表2b の注（lab_log 追記113）: 最も大きく折れ返るのは脈波伝播速度の行で、平均血圧の行も小さく折れ返る。
    # 「〜の行のみ単調でない」とは書かない。
    out["表2b"] = Cell({"ja": f"脈波伝播速度の行が最も大きく折れ返る: ΔT {v[0]} → {v[1]} → {v[2]} ms、RI {v[3]} → {v[4]} → {v[5]}",
                       "en": f"The PWV row reverses most: ΔT {v[0]} → {v[1]} → {v[2]} ms, RI {v[3]} → {v[4]} → {v[5]}"},
                      nums=[r2b[k]["value"] for k in keys])
    r2c = {r["id"]: r for r in t["表2c"]["rows"]}
    post2c = "\n".join(t["表2c"]["postscript"])
    m1 = re.search(r"ρ（大動脈PWV、起始部→指尖の伝播時間）= ([−+\d.]+)（(\d)/(\d)）", post2c)
    m2 = re.search(r"ρ（分解法 ΔT、同）= ([−+\d.]+)（(\d)/(\d)）", post2c)
    if not (m1 and m2):
        raise ValueError("表2c の注から ρ を取り出せない")
    a, rr = r2c["aorta_to_finger"], r2c["radial_to_finger"]
    out["表2c"] = Cell(
        {"ja": f"大動脈起始部→指尖 {a['median']['raw']}（{a['p5_p95']['raw']}）、橈骨→指尖 {rr['median']['raw']}（{rr['p5_p95']['raw']}）。"
               f"ρ(大動脈PWV, 起始部→指尖) {m1.group(1)}（{m1.group(2)}/{m1.group(3)}）、ρ(分解法 ΔT, 同) {m2.group(1)}（{m2.group(2)}/{m2.group(3)}）",
         "en": f"Aortic root→finger {a['median']['raw']} ({e(a['p5_p95']['raw'])}), radial→finger {rr['median']['raw']} ({e(rr['p5_p95']['raw'])}). "
               f"ρ(aortic PWV, root→finger) {m1.group(1)} ({m1.group(2)}/{m1.group(3)}); ρ(decomposition ΔT, same) {m2.group(1)} ({m2.group(2)}/{m2.group(3)})"},
        nums=[a["median"]["value"], a["p5_p95"]["lo"], a["p5_p95"]["hi"], rr["median"]["value"], rr["p5_p95"]["lo"], rr["p5_p95"]["hi"],
              float(m1.group(1).replace("−", "-")), float(m1.group(2)), float(m1.group(3)),
              float(m2.group(1).replace("−", "-")), float(m2.group(2)), float(m2.group(3))])
    return out


def build_table1(N: dict, P: dict) -> Table:
    items = P["prespecified_items"]
    results = t1_results(N)
    thr = N["meta"]["criterion"]["min_median_abs_rho"]
    k = N["meta"]["criterion"]["strata_required"]
    T = Table(
        "table1", {"ja": "表1", "en": "Table 1"},
        {"ja": "事前に決めた解析項目の一覧 ― 規則・凍結した日・結果・判定・出典{src}。"
               f"判定規準（実行前に凍結）: 年齢層内 Spearman ρ が全 {k} 層で予測の向き、かつ中央値 |ρ| ≥ {thr:.2f}。",
         "en": "Prespecified analysis items: rule, date frozen, result, verdict and source{src}. "
               f"Criterion frozen before the run: within-stratum Spearman ρ with the predicted sign in all {k} strata and median |ρ| ≥ {thr:.2f}."},
        {"ja": ["項目", "役割", "規則（凍結した日）", "結果", "判定{tier}", "出典"],
         "en": ["Item", "Role", "Rule (date frozen)", "Result", "Verdict{tier}", "Source"]},
        notes=[
            ("src", {"ja": "規則の原文は lab_log 2026-09-03「研究0 の事前規準と、実行前に捕まえた列ずれ」、`docs/research/gate0_rules_v2.md`"
                           "（2026-09-04）、lab_log 2026-09-04（追記）「決定試験の前に閾値を凍結した」。判定は `docs/research/roadmap_v1.md` §9 と同一。"
                           "結果の数値は `02_tables.md` 表1・表2・表2b・表2c（`data/paper2_numbers.json`）から取った。"
                           "「追記 n」は lab_log の追記番号。乱数の不変条件（28番）は実装の検査であり解析項目ではないので表に含めない。",
                     "en": "Rule texts from lab_log 2026-09-03 (prespecified criteria of study 0), `docs/research/gate0_rules_v2.md` (2026-09-04) "
                           "and lab_log 2026-09-04 (thresholds frozen before the decision test). Verdicts are identical to `docs/research/roadmap_v1.md` §9. "
                           "Result values are read from tables 1, 2, 2b and 2c of `02_tables.md` (`data/paper2_numbers.json`). "
                           "'Entry n' is a numbered addendum of lab_log. The random-beat invariants (script 28) are an implementation check, not an analysis item, and are omitted."}),
            ("tier", STAGE_NOTE),
        ],
        landscape=True, small=True)
    # prespec_chronology.json の項目名は 02_tables の表番号で始まる（表2・表2b・表2c）。この集では表3・3b・3c なので
    # 項目名の先頭だけを読み替える（役割の列の「表2 行n」は既にこの集の番号。出典の列は 02_tables の番号のまま）。
    name_ja_prefix = {"表2 振った": ("表2 ", "表3 "), "表2b": ("表2b ", "表3b "), "表2c": ("表2c ", "表3c ")}
    for key in T1_ORDER:
        it = next(i for i in items if key in i["name_ja"])
        en = T1_EN[key]
        rule_ja = f"{it['rule']}〔凍結 {it['frozen_on']}〕"
        rule_en = f"{en['rule']} [frozen {en['frozen']}]"
        name_ja = it["name_ja"]
        if key in name_ja_prefix and name_ja.startswith(name_ja_prefix[key][0]):
            name_ja = name_ja_prefix[key][1] + name_ja[len(name_ja_prefix[key][0]):]
        T.add(
            label_cell(name_ja, en["name"]),
            text_cell(T1_ROLE_JA[key], en["role"]),
            text_cell(rule_ja, rule_en, src=rule_ja),
            results[key],
            text_cell(T1_VERDICT_JA[key], en["verdict"]),
            text_cell(it["source"].replace("; 図3", ""), en["source"], src=it["source"]),
        )
    return T


# ---------------------------------------------------------------------------
# 表2  判定表（02 表1）
# ---------------------------------------------------------------------------

def build_table2(N: dict) -> Table:
    t = N["tables"]["表1"]
    meta = N["meta"]
    crit = meta["criterion"]
    n = fmt_int(meta["n_subjects_decision_test"])
    k = crit["strata_required"]
    hdr_ja = [c["header_ja"] for c in t["columns"]]
    T = Table(
        "table2", {"ja": "表2", "en": "Table 2"},
        {"ja": f"年齢層内 Spearman 順位相関の中央値（決定試験・{n} 名・{k} 層）{{src}}。判定規準は実行前に凍結: {crit['text']}。"
               f"{crit['sign_prediction']}表には絶対値を示し、括弧内は予測の符号を持った層の数{{bold}}。",
         "en": f"Median within-age-stratum Spearman rank correlation (decision test, {n} subjects, {k} strata){{src}}. "
               f"Criterion frozen before the run: predicted sign in all {k} strata and median |ρ| ≥ {crit['min_median_abs_rho']:.2f} "
               "(ΔT × aortic PWV negative, RI × peripheral vascular resistance positive). Magnitudes are shown; "
               "the number in parentheses is the number of strata with the predicted sign{bold}."},
        {"ja": hdr_ja, "en": ["Index construction", "ΔT × aortic PWV", "RI × peripheral vascular resistance", "Verdict"]},
        notes=[
            ("src", {"ja": "出典: `02_tables.md` 表1（`docs/research/roadmap_v1.md` §9、lab_log 追記12）。凍結版は 20番 `20_pwdb_validity.py`、"
                           "特徴点法と陽性対照は 23番 `23_pwdb_landmarks.py`、第2版と早期振幅比は 26番 `26_pwdb_compare.py`・27番 `27_threshold_sensitivity.py`"
                           "（`data/pwdb/pwdb_compare.csv`・`pwdb_compare_report.txt`）。",
                     "en": "Source: table 1 of `02_tables.md` (`docs/research/roadmap_v1.md` §9; lab_log entry 12). Frozen version: script 20 `20_pwdb_validity.py`; "
                           "fiducial-point analysis and positive control: script 23 `23_pwdb_landmarks.py`; rebuilt version and early amplitude ratio: "
                           "scripts 26 `26_pwdb_compare.py` and 27 `27_threshold_sensitivity.py` (`data/pwdb/pwdb_compare.csv`, `pwdb_compare_report.txt`)."}),
            ("bold", {"ja": f"太字は規準（{crit['bold_rule']}）を満たす値。「未検証」は評価していない指標、「—」は該当する指標が無いことを表す。",
                      "en": f"Bold, criterion met (median |ρ| ≥ {crit['min_median_abs_rho']:.2f} with the predicted sign in every stratum). "
                            "'Not evaluated', the index was not evaluated; —, no such index."}),
            ("tier", {"ja": "第2版の 2 行は採択数が少なく段で分けて示す。" + STAGE_NOTE["ja"],
                      "en": "The two rebuilt-version rows accepted few beats and are given by tier. " + STAGE_NOTE["en"]}),
        ])
    for r in t["rows"]:
        lab = LABELS["methods"][r["id"]]
        marker = "{tier}" if r["id"].startswith("pda_v2") else ""
        v = r["verdict"]
        T.add(
            label_cell(lab["ja"] + marker, lab["en"] + marker, bold=r["label_bold"]),
            rho_cell(r["dt_pwv"]),
            rho_cell(r["ri_pvr"]),
            Cell({"ja": v["raw"], "en": ("**" if v["bold"] else "") + VERDICT_EN[v["verdict"]] + ("**" if v["bold"] else "")},
                 nums=[], bold=False, kind="verdict"),
        )
    return T


# ---------------------------------------------------------------------------
# 表3・3b・3c  因子の主効果（02 表2・2b・2c）
# ---------------------------------------------------------------------------

def build_table3(N: dict) -> list[Table]:
    t2, t2b, t2c = N["tables"]["表2"], N["tables"]["表2b"], N["tables"]["表2c"]
    I = LABELS["indices"]
    F = LABELS["factors"]
    post = "\n".join(t2["postscript"])
    ratio_lm = find_num(post, r"脈波伝播速度の列が心拍数の列の (\d+) 倍", 1)
    ratio_pda = find_num(post, r"分解法の ΔT では ([\d.]+) 倍", 1)
    r2 = {r["id"]: r for r in t2["rows"]}
    rho_hr = r2["heart_rate"]["pda_dt"]["raw"]
    rho_pwv = r2["pwv"]["pda_dt"]["raw"]
    T = Table(
        "table3", {"ja": "表3", "en": "Table 3"},
        {"ja": "振った因子ごとの年齢層内主効果（+1SD と −1SD の平均の差 ÷ 層平均、%）{src}。括弧内はその因子と指標の年齢層内 Spearman 順位相関{bold}。",
         "en": "Within-stratum main effect of each varied factor (difference between the +1 SD and −1 SD means ÷ stratum mean, %){src}. "
               "In parentheses, the within-stratum Spearman rank correlation of the factor with the index{bold}."},
        {"ja": [c["header_ja"] for c in t2["columns"]],
         "en": ["Factor", I["landmark_dt"]["en"], I["pda_dt"]["en"], I["pda_ri"]["en"], I["amp_ratio"]["en"]]},
        notes=[
            ("src", {"ja": "出典: `02_tables.md` 表2（`docs/research/roadmap_v1.md` §9、lab_log 追記12。20番・23番・26番の因子別の表）。",
                     "en": "Source: table 2 of `02_tables.md` (`docs/research/roadmap_v1.md` §9; lab_log entry 12; factor tables of scripts 20, 23 and 26)."}),
            ("bold", {"ja": f"太字は本文で対比する行・列（脈波伝播速度と心拍数）。特徴点法の ΔT では脈波伝播速度の列が心拍数の列の {ratio_lm} 倍であるのに対し、"
                            f"分解法の ΔT では {ratio_pda} 倍にとどまる。順位では分解法の ΔT は心拍数（{unbold(rho_hr).split('（')[1].rstrip('）')}）が"
                            f"脈波伝播速度（{unbold(rho_pwv).split('（')[1].rstrip('）')}）を上回る。",
                      "en": f"Bold, the rows and columns contrasted in the text (pulse wave velocity and heart rate). For fiducial-point ΔT the pulse-wave-velocity "
                            f"column is {ratio_lm} times the heart-rate column; for decomposition ΔT only {ratio_pda} times. In rank terms decomposition ΔT follows "
                            f"heart rate ({unbold(rho_hr).split('（')[1].rstrip('）')}) more than pulse wave velocity ({unbold(rho_pwv).split('（')[1].rstrip('）')})."}),
        ])
    for r in t2["rows"]:
        cells = [label_cell(F[r["id"]]["ja"], F[r["id"]]["en"], bold=r["label_bold"])]
        for key in ("landmark_dt", "pda_dt", "pda_ri", "amp_ratio"):
            c = r[key]
            nums = [c["effect_pct"]] + ([c["rho"]] if c.get("rho") is not None else [])
            cells.append(Cell({"ja": c["raw"], "en": en_from_raw(c["raw"])}, nums=nums, bold=c["bold"]))
        T.add(*cells)

    # 02 表2b の注（lab_log 追記113）: 最も大きく折れ返るのは脈波伝播速度の行、ΔT では平均血圧の行も小さく折れ返る。
    r2b_map = {r["id"]: r for r in t2b["rows"]}["map"]
    map_dt = [unbold(r2b_map[k]["raw"]) for k in ("dt_minus1sd", "dt_base", "dt_plus1sd")]
    post_2b = t2b["postscript"][0]
    lead_2b = "脈波伝播速度の行が最も大きく折れ返る。"       # 注の冒頭と重複するので、その文だけ省く（文が変われば全文を載せる）
    if post_2b.startswith(lead_2b):
        post_2b = post_2b[len(lead_2b):]
    Tb = Table(
        "table3b", {"ja": "表3b", "en": "Table 3b"},
        {"ja": "1 因子掃引（他の因子は基準値。年齢層の中央値。ΔT は ms）{src}{note}。",
         "en": "One-factor-at-a-time sweep (other factors at their reference values; median over age strata; ΔT in ms){src}{note}."},
        {"ja": [c["header_ja"] for c in t2b["columns"]],
         "en": ["Factor", "ΔT, −1 SD", "ΔT, reference", "ΔT, +1 SD", "RI, −1 SD", "RI, reference", "RI, +1 SD"]},
        notes=[
            ("src", {"ja": "出典: `02_tables.md` 表2b（lab_log 2026-09-03「研究0 の結果」・追記12。20番）。分解法・凍結版の値。",
                     "en": "Source: table 2b of `02_tables.md` (lab_log 2026-09-03, results of study 0; entry 12; script 20). Values of the frozen decomposition."}),
            ("note", {"ja": "太字は最も大きく折れ返る行（脈波伝播速度）。" + post_2b,
                      "en": "Bold, the row with the largest reversal (pulse wave velocity): ΔT moves against the prediction on the −1 SD side "
                            f"and RI is U-shaped. In ΔT the mean-arterial-pressure row also reverses slightly ({map_dt[0]} → {map_dt[1]} → {map_dt[2]}; "
                            "lab_log entry 113)."}),
        ])
    for r in t2b["rows"]:
        cells = [label_cell(F[r["id"]]["ja"], F[r["id"]]["en"], bold=r["label_bold"])]
        for c in t2b["columns"][1:]:
            cc = r[c["key"]]
            cells.append(Cell({"ja": cc["raw"], "en": unbold(cc["raw"]) if not cc["bold"] else cc["raw"]}, nums=[cc["value"]], bold=cc["bold"]))
        Tb.add(*cells)

    post_c = "\n".join(t2c["postscript"])
    m1 = re.search(r"ρ（大動脈PWV、起始部→指尖の伝播時間）= ([−+\d.]+)（(\d)/(\d)）", post_c)
    m2 = re.search(r"ρ（分解法 ΔT、同）= ([−+\d.]+)（(\d)/(\d)）", post_c)
    Tc = Table(
        "table3c", {"ja": "表3c", "en": "Table 3c"},
        {"ja": "真の伝播時間（モデルの立ち上がり時刻表より）{src}{note}。",
         "en": "True transit times from the model onset-time table{src}{note}."},
        {"ja": [c["header_ja"] for c in t2c["columns"]], "en": ["Segment", "Median", "5th–95th percentile"]},
        notes=[
            ("src", {"ja": "出典: `02_tables.md` 表2c（lab_log 2026-09-03「研究0 の結果」。20番。PWDB 配布物 `pwdb_onset_times.csv`）。",
                     "en": "Source: table 2c of `02_tables.md` (lab_log 2026-09-03, results of study 0; script 20; PWDB file `pwdb_onset_times.csv`)."}),
            ("note", {"ja": post_c.replace("\n", ""),
                      "en": f"Within-stratum ρ(aortic PWV, root→finger transit time) = {m1.group(1)} ({m1.group(2)}/{m1.group(3)}); "
                            f"ρ(decomposition ΔT, same) = {m2.group(1)} ({m2.group(2)}/{m2.group(3)})."}),
        ])
    seg_en = {"aorta_to_finger": "Aortic root → finger", "radial_to_finger": "Radial → finger"}
    for r in t2c["rows"]:
        Tc.add(label_cell(r["label_ja"], seg_en[r["id"]]),
               Cell({"ja": r["median"]["raw"], "en": r["median"]["raw"]}, nums=[r["median"]["value"]]),
               Cell({"ja": r["p5_p95"]["raw"], "en": en_from_raw(r["p5_p95"]["raw"])}, nums=[r["p5_p95"]["lo"], r["p5_p95"]["hi"]]))
    return [T, Tb, Tc]


# ---------------------------------------------------------------------------
# 表4  基底関数と成分数（02 表3。探索・事後）
# ---------------------------------------------------------------------------

BASIS_EN = {
    "skewgauss_a08": "Skew-Gaussian, α∈[0,8] (frozen)",
    "gauss": "Gaussian (no skew)",
    "skewgauss_pm8": "Skew-Gaussian, α∈[−8,8] (Basso)",
    "gamma_frozen": "Gamma (frozen search bounds)",
    "gamma_wide": "Gamma (wide search bounds)",
}


def after_bold(text: str) -> str:
    """preamble の先頭の太字の文（探索・事後の断り）を除いた残りを 1 行にする。"""
    s = re.sub(r"^\*\*[^*]+\*\*\s*", "", text.strip())
    return re.sub(r"\s*\n\s*", "", s)


def build_table4(N: dict) -> Table:
    t = N["tables"]["表3"]
    pre = "\n".join(t["preamble"])
    m = re.search(r"対象は (\d+) 名から取った切痕のある (\d+) 拍で、層は (\d+) 拍以上の (\d+) 層", pre)
    m2 = re.search(r"決定試験の ([\d,]+) 名・(\d) 層より小さく", pre)
    if not (m and m2):
        raise ValueError("表3 の前書きから対象の数を取り出せない")
    T = Table(
        "table4", {"ja": "表4", "en": "Table 4"},
        {"ja": "基底関数と成分数の総当たり（探索・事後）{posthoc}{src}{subj}{bold}。",
         "en": "Basis function and number of components, exhaustive sweep (exploratory, post hoc){posthoc}{src}{subj}{bold}."},
        {"ja": [c["header_ja"] for c in t["columns"]],
         "en": ["Basis", "Components", "Wang pass", "Errx median [ms]", "NRMSE median", "Parameters on a bound", "ΔT × PWV |ρ|"]},
        notes=posthoc_notes("表3") + [
            ("src", source_note(t, "台本 31番 `31_pwdb_basis_explore.py`（`data/pwdb/pwdb_basis_explore.csv`）。",
                                "Script 31 `31_pwdb_basis_explore.py` (`data/pwdb/pwdb_basis_explore.csv`).")),
            ("subj", {"ja": after_bold(pre),
                      "en": f"Subjects: {m.group(2)} notched beats taken from {m.group(1)} subjects, in {m.group(4)} strata of at least "
                            f"{m.group(3)} beats (smaller than the {m2.group(1)} subjects and {m2.group(2)} strata of the decision test, so the "
                            "values are wider). Preprocessing and Wang's acceptance criterion identical to the decision test. |ρ| is the "
                            "within-stratum median over all beats (acceptance ignored); in parentheses, the number of strata with the predicted sign."}),
            ("bold", {"ja": "太字は分解の中で最も追った 2 行。", "en": "Bold, the two best-tracking rows among the decompositions."}),
        ])
    for r in t["rows"]:
        cells = [label_cell(r["label_ja"], BASIS_EN[r["basis_id"]], bold=r["label_bold"])]
        for key in ("n_components", "wang_pass", "errx_median_ms", "nrmse_median", "boundary_fraction"):
            cells.append(value_cell(r[key]))
        cells.append(rho_cell(r["dt_pwv"]))
        T.add(*cells)
    return T


# ---------------------------------------------------------------------------
# 表5a・5b  公表された条件の再現と逸脱表（02 表4）
# ---------------------------------------------------------------------------

T5_EN = {
    "tigges2017": "Tigges 2017 (recommended gamma, M = 3)",
    "fleischhauer2020": "Fleischhauer 2020 (two kernels)",
    "couceiro2015": "Couceiro 2015",
    "wang2013": "Wang 2013 (all beats)",
    "basso2024": "Basso 2024 (L = 3)",
    "goswami2010": "Goswami 2010",
    "landmark_ref": "Reference: fiducial-point indices supplied with PWDB (Charlton)",
}
DEV_LABEL_EN = {"all": "All conditions", "wang2013": "Wang 2013", "couceiro2015": "Couceiro 2015",
                "tigges2017": "Tigges 2017", "common": "Common"}
DEV_TEXT_EN = {
    "all": "The optimiser was that of our environment (SciPy), not identical to the originals",
    "wang2013": "The weight search step was set to 1",
    "couceiro2015": "For the 1% of beats whose fifth-component initial value could not be formed from table 1 of the original, the default was used",
    "tigges2017": "Resampling to 40 Hz for the corrected Akaike information criterion used a Kaiser window, as the original describes",
    "common": "PWDB gives one noise-free beat per subject; the conditions differ from the measured waveforms of the originals",
}


def build_table5(N: dict) -> list[Table]:
    t, td = N["tables"]["表4"], N["tables"]["表4_逸脱表"]
    pre = "\n".join(t["preamble"])
    m = re.search(r"系統的に選んだ ([\d,]+) 名（subj_no を (\d) で割った剰余が (\d)）", pre)
    if not m:
        raise ValueError("表4 の前書きから対象の数を取り出せない")
    rows = {r["id"]: r for r in t["rows"]}
    ref_dt, ref_ri = rows["landmark_ref"]["dt_pwv"]["raw"], rows["landmark_ref"]["ri_pvr"]["raw"]
    cou_ri = unbold(rows["couceiro2015"]["ri_pvr"]["raw"])
    T = Table(
        "table5a", {"ja": "表5a", "en": "Table 5a"},
        {"ja": "公表された条件をそのまま当てた結果（探索・事後）{posthoc}{src}{subj}{bold}{ri}。",
         "en": "Published fitting conditions applied unchanged (exploratory, post hoc){posthoc}{src}{subj}{bold}{ri}."},
        {"ja": [c["header_ja"] for c in t["columns"]],
         "en": ["Source of the conditions", "ΔT × aortic PWV", "RI × peripheral vascular resistance"]},
        notes=posthoc_notes("表4") + [
            ("src", source_note(t, "台本 33番 `33_pwdb_literature_replica.py`（`docs/research/results/33_literature_replica_report.txt`、lab_log 追記16〜18）。",
                                "Script 33 `33_pwdb_literature_replica.py` (`docs/research/results/33_literature_replica_report.txt`; lab_log entries 16–18).")),
            ("subj", {"ja": after_bold(pre),
                      "en": f"Subjects: {m.group(1)} chosen systematically (subject number modulo {m.group(2)} equal to {m.group(3)}). In the same "
                            f"subset the fiducial-point indices supplied with PWDB (Charlton) give ΔT {ref_dt} and RI {ref_ri}. "
                            "In parentheses, the number of strata with the predicted sign."}),
            ("bold", {"ja": "太字は ΔT で最良の条件と、RI で特徴点由来を上回った唯一の条件。",
                      "en": "Bold, the best condition for ΔT and the only condition whose RI exceeded the fiducial-point value."}),
            ("ri", {"ja": re.sub(r"\s*\n\s*", "", unbold(t["postscript"][1])),
                    "en": f"For RI only Couceiro's R1_d exceeds the fiducial-point value ({cou_ri.replace('R1_d ', '').split('（')[0]} versus {ref_ri}). "
                          "Under the prespecified rule this row is not taken into the verdict. Three caveats: in this model peripheral vascular "
                          "resistance determines the diastolic decay, so part of the association is built in; measured waveforms often lack a "
                          "diastolic peak; and deviations remain in our replication."}),
        ])
    for r in t["rows"]:
        T.add(label_cell(r["label_ja"], T5_EN[r["id"]], bold=r["label_bold"]), rho_cell(r["dt_pwv"]), rho_cell(r["ri_pvr"]))

    Td = Table(
        "table5b", {"ja": "表5b", "en": "Table 5b"},
        {"ja": "逸脱表 ― 文献の条件の再現で原著と異なる点（表5a の補足）{posthoc}{src}。",
         "en": "Deviations from the published conditions in the replication (supplement to table 5a){posthoc}{src}."},
        {"ja": [c["header_ja"] for c in td["columns"]], "en": ["Source", "Deviation"]},
        notes=posthoc_notes("表4") + [
            ("src", source_note(td, "台本 33番の逸脱表（lab_log 追記17・18）。", "Deviation table of script 33 (lab_log entries 17, 18).")),
        ])
    for r in td["rows"]:
        raw = r["deviation_ja"]["text"]
        Td.add(label_cell(r["label_ja"], DEV_LABEL_EN[r["id"]]), text_cell(raw, DEV_TEXT_EN[r["id"]], src=raw))
    return [T, Td]


# ---------------------------------------------------------------------------
# 表6a・6b  波形の型で分けた関連と、真値の層内の幅（02 表5・5b）
# ---------------------------------------------------------------------------

def type_label(r: dict, lang: str) -> str:
    tid = r["id"]
    if tid == "all":
        return r["label_ja"] if lang == "ja" else f"{LABELS['types']['all']['en']} ({fmt_int(r['n'])})"
    if lang == "ja":
        return r["label_ja"]
    base = LABELS["types"][tid]["en"]
    if r.get("strata_count"):
        return f"{base} ({fmt_int(r['n'])}; {r['strata_count']} strata)"
    return f"{base} ({fmt_int(r['n'])})"


def build_table6(N: dict) -> list[Table]:
    t, tb = N["tables"]["表5"], N["tables"]["表5b"]
    n_all = fmt_int(N["meta"]["n_subjects_decision_test"])
    pre = after_bold(t["preamble"][0])
    T = Table(
        "table6a", {"ja": "表6a", "en": "Table 6a"},
        {"ja": "波形の型で分けた年齢層内 Spearman 順位相関（探索・事後）{posthoc}{src}{tier}{defs}{bold}。",
         "en": "Within-age-stratum Spearman rank correlation by waveform type (exploratory, post hoc){posthoc}{src}{tier}{defs}{bold}."},
        {"ja": [c["header_ja"] for c in t["columns"]],
         "en": ["Waveform type (n)", "Fiducial-point ΔT, tier C", "Decomposition ΔT (frozen), tier A", "Early amplitude ratio, tier C",
                "Fiducial-point RI, tier C", "Decomposition RI (frozen), tier A"]},
        notes=posthoc_notes("表5") + [
            ("src", source_note(t, "50番 節C が同じ入力から型1・型3 の特徴点法の値を再現している。",
                                "Section C of script 50 reproduces the fiducial-point values of types 1 and 3 from the same input.")),
            ("tier", {"ja": STAGE_NOTE["ja"] + "。特徴点由来の指標は採否の判定を持たないので常に C 段。",
                      "en": STAGE_NOTE["en"] + ". Fiducial-point indices carry no acceptance decision and are therefore always tier C."}),
            ("defs", {"ja": pre + t["postscript"][0],
                      "en": f"Subjects as in the decision test ({n_all}); waveform type by `pda2.find_landmarks` (type 1, dicrotic notch and "
                            "diastolic peak present as extrema; type 3, no extrema, but the inflection point of the descending limb can stand in; "
                            "type 4, neither found). Six age strata, at least 8 subjects per stratum. Median |ρ|; in parentheses, the number of "
                            "strata with the predicted sign. The three left columns are ΔT × aortic PWV (predicted sign negative; positive for "
                            "the early amplitude ratio), the two right columns RI × peripheral vascular resistance (positive)."}),
            ("bold", {"ja": f"太字は規準（{N['meta']['criterion']['bold_rule']}）を満たす値。",
                      "en": f"Bold, criterion met (median |ρ| ≥ {N['meta']['criterion']['min_median_abs_rho']:.2f} with the predicted sign in every stratum)."}),
        ])
    for r in t["rows"]:
        digit = int(r["id"][-1]) if r["id"] != "all" else None
        nums = ([float(digit)] if digit else []) + [float(r["n"])] + ([float(r["strata_count"])] if r.get("strata_count") else [])
        cells = [label_cell(type_label(r, "ja"), type_label(r, "en"), nums=nums)]
        for c in t["columns"][1:]:
            cells.append(rho_cell(r[c["key"]]))
        T.add(*cells)

    pre_b = after_bold(tb["preamble"][0])
    Tb = Table(
        "table6b", {"ja": "表6b", "en": "Table 6b"},
        {"ja": "真値そのものの層内のばらつき（範囲の制限。探索・事後・C 段）{posthoc}{src}{tier}{defs}。",
         "en": "Within-stratum spread of the true values themselves (restriction of range; exploratory, post hoc, tier C){posthoc}{src}{tier}{defs}."},
        {"ja": [c["header_ja"] for c in tb["columns"]],
         "en": ["Waveform type", "Width of aortic PWV [m/s]", "Width of peripheral vascular resistance", "Strata with ≥ 8 subjects", "Smallest stratum n"]},
        notes=posthoc_notes("表5b") + [
            ("src", source_note(tb)),
            ("tier", STAGE_NOTE),
            ("defs", {"ja": "型1 の値を「手法の失敗」と読んではいけない。型1 に絞ると年齢層の中で大動脈脈波伝播速度がほとんど動かなくなるためで、"
                            "表6a の型1 の列はどの手法でも低い。四分位範囲の幅を年齢層ごとに出し、層をまたいだ中央値を示す（数値は PWDB の配布物の単位のまま）。"
                            + tb["postscript"][1] + "太字は本文で対比する値。",
                      "en": "Type-1 values must not be read as a failure of the methods: restricted to type 1, aortic PWV barely varies within an age "
                            "stratum. The interquartile range was taken per age stratum and its median across strata is shown (units as distributed "
                            "with PWDB). This table was made after the weakness of type 1 was seen; no prediction was fixed and no test is made. "
                            "Bold, the values contrasted in the text."}),
        ])
    lab_en = {"type1": LABELS["types"]["type1"]["en"], "type3": LABELS["types"]["type3"]["en"],
              "type4": LABELS["types"]["type4"]["en"], "ratio_type1_type3": "Type 1 ÷ type 3"}
    for r in tb["rows"]:
        s8 = r["strata_ge8"]
        Tb.add(label_cell(r["label_ja"], lab_en[r["id"]], bold=r["label_bold"]),
               value_cell(r["pwv_iqr_width_m_s"]),
               Cell({"ja": r["pvr_iqr_width"]["raw"], "en": r["pvr_iqr_width"]["raw"]}, nums=[r["pvr_iqr_width"]["value"]],
                    bold=r["pvr_iqr_width"]["bold"]),
               Cell({"ja": s8["raw"], "en": s8["raw"]}, nums=[] if s8.get("dash") else [float(s8["num"]), float(s8["den"])]),
               value_cell(r["min_stratum_n"]))
    return [T, Tb]


# ---------------------------------------------------------------------------
# 表7  診断と改善案の対応（prespec_chronology.json variants_table6 ＋ 02 表6 の数値）
# ---------------------------------------------------------------------------

VARIANT_MAP = {   # 02 表6 の行 id → variants_table6 の id
    "fb": "fb (0)", "dmu001": "relax (2)", "reservoir": "conv (4)", "reservoir_tau015": "conv15 (4c)",
    "decay": "decay (3)", "conv": "convout (7)", "twostage": "deconv2 (8)", "trunc065": "trunc08 (6b)",
    "trunc055": "trunc055 (6c)", "trunc075": "trunc075 (6d)", "trunc_abs045": "truncabs (6e)",
    "deriv": "deriv (9)", "landmark": "lm_supplied（参考）",
}

# 表7 の文章の列（和文 60 字以内・英文 20 語以内。lab_log の追記と照合済み）
T7_TEXT = {
    "fb": dict(
        what=("何も変えない。src/pda.py の fit_beat を既定の引数で呼ぶ（Δμ の下限 0.08 s）",
              "Nothing changed: fit_beat of src/pda.py with its default arguments (delay bound 0.08 s)"),
        why=("対照。複製が本体と同じ振る舞いをすることの照合（追記142）",
             "Control; checks that the replica behaves exactly like the original (entry 142)"),
        pred=("P10: ΔT・RI・採否が 26番の列と全例一致し、同梱の特徴点の型別 ρ が 48番を再現（追記143）",
              "P10: ΔT, RI and acceptance match script 26 for all subjects; fiducial-point ρ by type reproduces script 48 (entry 143)"),
        read=("全例で 26番の列と一致。節C が読む拍・真値・規約は 26番と同じ（追記144）",
              "Identical to script 26 for every subject: section C reads the same beats, truth and rules (entry 144)")),
    "dmu001": dict(
        what=("探索範囲の Δμ の下限だけを 0.08 → 0.01 s に緩める。ほかは凍結版のまま",
              "Only the lower bound of Δμ relaxed from 0.08 to 0.01 s; otherwise the frozen fit"),
        why=("合成脈波で凍結版の ΔT が真値 150 ms 付近で下限に詰まり、RI の符号が反転した（追記141・143）",
             "In synthetic beats frozen ΔT piled up at the bound near 150 ms and RI reversed sign (entries 141, 143)"),
        pred=("P7: 型3・C 段で Δμ 下限に張り付く割合が (4) より 0.10 以上大きい（追記143）",
              "P7: type 3, tier C: fraction at the Δμ bound exceeds (4) by ≥ 0.10 (entry 143)"),
        read=("下限に張り付いた拍は全型で皆無。Δμ の下限は原因ではない（追記144）",
              "No beat in any type sits at the bound; the Δμ bound is not the cause (entry 144)")),
    "reservoir": dict(
        what=("前進波に単位面積の指数核を畳み込んだ貯留槽の項を足す（母数 g・τ を追加、τ ≥ 0.05 s）",
              "Reservoir term added: forward wave convolved with a unit-area exponential kernel (parameters g, τ; τ ≥ 0.05 s)"),
        why=("追記139 の候補 (1)。2 要素 Windkessel の解。合成では ρ 1.00・縮退なし（追記143）",
             "Candidate (1) of entry 139; two-element Windkessel solution. In synthetic beats ρ 1.00, no degeneracy, accepted (entry 143)"),
        pred=("P6・P8・P9: 型3 の A 段で ΔT・RI（75 歳層も正）が規準、型1 でも ΔT が規準（追記143）",
              "P6, P8, P9: type-3 tier-A ΔT meets criterion; RI positive even at 75 years; type-1 ΔT too (entry 143)"),
        read=("P6・P8・P9 はすべて「いいえ」。g が上限・τ が下限に達し、第1成分の定数倍に縮退（追記144・146）",
              "All three 'no'; g at its upper and τ at its lower bound, i.e. degenerate (entries 144, 146)")),
    "reservoir_tau015": dict(
        what=("(4) と貯留槽の時定数 τ の下限だけが違う（0.05 → 0.15 s）",
              "Same as (4) except the lower bound of the reservoir time constant τ (0.05 → 0.15 s)"),
        why=("追記146 の縮退の診断（τ → 0 で核が δ 関数に近づく）。0.15 s は生理の範囲で前進波の幅の 3 倍",
             "Degeneracy diagnosis of entry 146 (kernel → δ as τ → 0); 0.15 s is physiological and 3× forward-wave width"),
        pred=("「結果を見る前に決めた値。縮退が外れても効かないなら、項を足す案は縮退を塞いでも効かないと書ける」（追記146）",
              "'Fixed before seeing the result; if it still fails, adding a term fails even when degeneracy is blocked' (entry 146)"),
        read=("張り付きが新しい下限へ移っただけで縮退は解消しない。通過率は全案で最低（追記147）",
              "The pile-up merely moved to the new bound; degeneracy persists and the pass rate is the lowest (entry 147)")),
    "decay": dict(
        what=("自由な指数減衰 d·exp(−(t−t0)/τ) を足す（24番 A1 と同じ形。10 母数）",
              "A free exponential decay d·exp(−(t−t0)/τ) added (same form as A1 of script 24; 10 parameters)"),
        # 追記139 の候補 (1)〜(4) に減衰項は無い（追記139 は 24番 A1 を「PWDB では未実行」と記録しただけ）
        why=("24番の減衰項（追記139 では未実行）。合成では効かず、実データでは下降が主因で順位が変わりうる（追記144 §4）",
             "Script-24 term, untested on PWDB (entry 139); failed in synthetic beats, but the decline dominates real data (entry 144 §4)"),
        pred=("数値の予測は無く「実データでどう出るかを見る」（追記144・145）",
              "No numerical prediction: 'see how it behaves in real data' (entries 144, 145)"),
        read=("凍結版よりわずかに良いだけ。τ が下限に達し (4) と同じ向きに縮退する（追記146）",
              "Only slightly better than the frozen fit; τ reaches its lower bound, degenerating like (4) (entry 146)")),
    "conv": dict(
        what=("2 成分の和に単位面積の指数核を掛けて置き換える（9 母数）。ΔT・RI は畳み込む前のピークから出す",
              "Two-component sum replaced by its convolution with a unit-area exponential kernel (9 parameters); ΔT, RI from unconvolved peaks"),
        why=("PWDB の PPG の生成（Charlton 2019 式 A1）に最も近く、(4) の縮退の向きが無い（追記149）",
             "Closest to the generation of the PWDB PPG (Charlton 2019, eq. A1), without the degenerate direction of (4) (entry 149)"),
        # 「直ると期待した」が記録されているのは追記152 §4（「追記149 では『(7) なら直る』と期待したが外れた」）
        pred=("「PWDB の PPG はこの形で作られ in silico では有利」と留保。直ると期待した（追記149・152）",
              "Caveat: 'the PWDB PPG is generated in this form, so favoured in silico'; expected to work (entries 149, 152)"),
        read=("波形には最もよく合うが ΔT の関連は凍結版より悪い。下降の自由度が第2成分の位置を不定にする（追記152）",
              "Fits waveform best yet ΔT tracks worse than frozen fit; freedom for the decline leaves component 2 undetermined (entry 152)")),
    "twostage": dict(
        what=("拍の後ろ 30% から τ を推定し、x = y + τ·y′ で逆畳み込みしてから凍結版を当てる",
              "τ estimated from the last 30% of the beat, deconvolution x = y + τ·y′, then the frozen fit"),
        why=("貯留槽の文献が拡張期後期から時定数を推定するやり方に対応（追記149）",
             "Corresponds to reservoir studies that estimate the time constant from late diastole (entry 149)"),
        pred=("「微分を使うので雑音に弱い」。数値の予測は無し（追記149）",
              "'Uses a derivative, so sensitive to noise'; no numerical prediction (entry 149)"),
        read=("逆畳み込みが行き過ぎて拡張期が持ち上がり、第2成分の幅が上限に達する。通過率は最低級（追記152）",
              "Deconvolution overshoots, lifting diastole; the width of component 2 hits its upper bound; pass rate near the lowest (entry 152)")),
    "trunc065": dict(
        what=("残差だけを拍長の 0.65 倍までで取る。探索範囲・起点・ピークは全長で計算。Δμ の下限は 0.08 s のまま",
              "Residual taken up to 0.65 of the beat only; bounds, starts, peaks use the full beat; Δμ bound unchanged"),
        why=("追記139 の候補 (3)。0.65 は 24番・追記139 で実データを見る前に置いた値（追記145）",
             "Candidate (3) of entry 139; 0.65 was set (script 24, entry 139) before real data were seen (entry 145)"),
        pred=("「主の割合は 0.65T のまま。見てから選び直さない」（追記145）。(6b) は (6) と同じはず（追記144）",
              "'Main fraction stays 0.65T; not re-chosen after seeing the numbers' (entry 145); (6b) should equal (6) (entry 144)"),
        read=("ΔT は A 段・C 段とも規準を満たし特徴点法を上回る。RI は A 段だけ。打ち切りだけが効く（追記146・147）",
              "ΔT meets the criterion in tiers A and C, above fiducial-point analysis; RI only in tier A (entries 146, 147)")),
    "trunc055": dict(
        what=("(6b) と打ち切りの割合だけが違う（0.55 倍。Δμ の下限は 0.08 s のまま）",
              "Differs from (6b) only in the truncation fraction (0.55; Δμ bound 0.08 s unchanged)"),
        why=("結論が割合の選び方に敏感でないことを示す記述のため（追記145 取り決め 2）",
             "Descriptive check that the conclusion is insensitive to the chosen fraction (entry 145, rule 2)"),
        pred=("「0.55 や 0.75 のほうが良く出ても主の割合は 0.65T と書く」（追記145 取り決め 3）",
              "'Even if 0.55 or 0.75 comes out better, the main fraction is written as 0.65T' (entry 145, rule 3)"),
        read=("C 段は 0.65 倍と同じ水準で全層の向きが合う。通過率は割合を振った中で最も低い（追記147）",
              "Tier C at the level of 0.65, all strata in direction; lowest pass rate among the fractions (entry 147)")),
    "trunc075": dict(
        what=("(6b) と打ち切りの割合だけが違う（0.75 倍）",
              "Differs from (6b) only in the truncation fraction (0.75)"),
        why=("同上（感度の記述。追記145 の取り決め）",
             "As above (descriptive sensitivity check; the rules of entry 145)"),
        pred=("同上（追記145 の取り決め）",
              "As above (the rules of entry 145)"),
        read=("C 段は同じ水準で向きが合い、通過率は割合を振った中で最も高い（追記147）",
              "Same level in tier C, in direction; the highest pass rate among the fractions (entry 147)")),
    "trunc_abs045": dict(
        what=("(6b) と切る位置の決め方だけが違う（絶対時間 0.45 s。拍長の 0.90 倍を上限）",
              "Differs from (6b) only in how the cut is placed: absolute 0.45 s (at most 0.90 of the beat)"),
        why=("W1: 拍長の割合で切ると切る位置が心拍数の関数になる（追記148・149）",
             "W1: cutting at a fraction of the beat makes the cut position a function of heart rate (entries 148, 149)"),
        pred=("「絶対時間で切っても同じ結論が出るかを見る」（追記149）",
              "'See whether cutting at an absolute time gives the same conclusion' (entry 149)"),
        read=("割合で切った版と同じ向き・同じ水準以上。打ち切りの効果は心拍数依存の産物ではない（追記152）",
              "Same direction, same or higher level than the fractional cut; not a product of heart-rate dependence (entry 152)")),
    "deriv": dict(
        what=("残差を d/dt(g1+g2) と y′ の差で取る。模型・探索範囲・起点・解の選び方は凍結版と同じ",
              "Residual taken between d/dt(g1+g2) and y′; model, bounds, starts and solution choice as in the frozen fit"),
        why=("追記139 の候補 (3) の後半。拡張期の下降は微分すると小さくなり、第2成分が引かれにくい（追記149）",
             "Candidate (3) of entry 139, second half: the decline shrinks under differentiation, pulling component 2 less (entry 149)"),
        pred=("「雑音 0.01 で通過率か型3・C 段の ΔT が 0.30 未満なら微分領域は雑音に弱い」（追記153 §5）",
              "'Pass rate < 0.30 or type-3 tier-C ΔT < 0.30 at 0.01 noise → derivative domain noise-sensitive' (entry 153 §5)"),
        read=("ΔT と RI とも C 段の規準を満たす唯一の版だが、雑音 2% で崩れ、採否が硬い側で落ちる（追記152・158）",
              "Only version meeting tier-C criterion for both ΔT and RI; degrades at 2% noise, rejects stiff subjects (entries 152, 158)")),
    "landmark": dict(
        what=("何も当てはめない。Charlton 2019 Table A3 の規則で同梱の特徴点から ΔT・RI を取る",
              "No fit; ΔT and RI taken from the fiducial points supplied with PWDB (Charlton 2019, table A3 rules)"),
        why=("決定試験の対照（2026-09-03・09-06）",
             "Control of the decision test (2026-09-03, 2026-09-06)"),
        pred=("決定試験で確定済み。48番の P1（型1・型3 のどちらでも規準）は型1 で外れた（追記137）",
              "Settled in the decision test; script-48 P1 (criterion in types 1 and 3) missed for type 1 (entry 137)"),
        read=("型3 の参照値。B 段は分解に有利な部分集合なので C 段より低い（追記158）",
              "Reference for type 3; tier B is a subset favourable to decomposition, hence below tier C (entry 158)")),
}


def variant_label(vid: str, lang: str, bold: bool = False) -> Cell:
    v = LABELS["variants"][vid]
    no = (v["no"] + " ") if v["no"] else ""
    return label_cell(no + v["ja"], no + v["en"], bold=bold)


def build_table7(N: dict, P: dict) -> Table:
    t6 = N["tables"]["表6"]
    variants = {v["id"]: v for v in P["variants_table6"]}
    T = Table(
        "table7", {"ja": "表7", "en": "Table 7"},
        {"ja": "診断と改善案の対応 ― 何を変え、どの観察から、走らせる前に何を予測し、どうなったか（型3・探索・事後）{posthoc}{src}{tier}{bold}。",
         "en": "Diagnosis and the corresponding modifications: what was changed, the observation behind it, the prediction fixed before the run, "
               "and the result (type 3; exploratory, post hoc){posthoc}{src}{tier}{bold}."},
        {"ja": ["改良案（番号）", "何を変えたか", "根拠となった観察（出典）", "走らせる前に固定した予測（出典）",
                "結果（型3: ΔT C 段・RI C 段・通過率）", "読み"],
         "en": ["Modification (no.)", "What was changed", "Observation behind it (source)", "Prediction fixed before the run (source)",
                "Result (type 3: ΔT tier C, RI tier C, pass rate)", "Reading"]},
        notes=posthoc_notes("表6") + [
            ("src", source_note(t6, "文章の列は `data/prespec_chronology.json`（lab_log の追記を起こしたもの）を 60 字以内に縮め、追記と照合した。"
                                    "数値の列は `02_tables.md` 表6（この集の表8）と同一。「追記 n」は lab_log の追記番号。",
                                "Text columns condense `data/prespec_chronology.json` (reconstructed from lab_log) to at most 20 words and were checked "
                                "against the entries cited; the numeric column is identical to table 6 of `02_tables.md` (table 8 of this set). "
                                "'Entry n' is a numbered addendum of lab_log.")),
            ("tier", STAGE_NOTE),
            ("bold", {"ja": f"太字は規準（{N['meta']['criterion']['bold_rule']}）を満たす値。",
                      "en": f"Bold, criterion met (median |ρ| ≥ {N['meta']['criterion']['min_median_abs_rho']:.2f} with the predicted sign in every stratum)."}),
        ],
        landscape=True, small=True)
    for r in t6["rows"]:
        v = variants[VARIANT_MAP[r["id"]]]
        src = " ".join(str(v[k]) for k in ("label_ja", "label_en", "what_changed", "why_proposed", "prediction", "outcome", "source"))
        tx = T7_TEXT[r["id"]]
        dc, rc_, pr = r["dt_pwv_C"], r["ri_pvr_C"], r["pass_rate"]
        res = Cell({"ja": f"ΔT C 段 {dc['raw']}・RI C 段 {rc_['raw']}・通過率 {pr['raw']}",
                    "en": f"ΔT tier C {en_from_raw(dc['raw'])}; RI tier C {en_from_raw(rc_['raw'])}; pass rate {pr['raw']}"},
                   nums=rho_cell(dc).nums + rho_cell(rc_).nums + value_cell(pr).nums)
        T.add(variant_label(r["id"], "ja"),
              text_cell(tx["what"][0], tx["what"][1], src=src),
              text_cell(tx["why"][0], tx["why"][1], src=src),
              text_cell(tx["pred"][0], tx["pred"][1], src=src),
              res,
              text_cell(tx["read"][0], tx["read"][1], src=src))
    return T


# ---------------------------------------------------------------------------
# 表8  拡張期の下降の扱いを変えたときの型3 の関連（02 表6）
# ---------------------------------------------------------------------------

GROUP_KEY = {"下降を説明する項を足す": "add", "下降を当てはめの対象から外す・残差の中で小さくする": "remove"}


def build_table8(N: dict) -> Table:
    t = N["tables"]["表6"]
    n_all = fmt_int(N["meta"]["n_subjects_decision_test"])
    post = "\n".join(t["postscript"])
    frac = find_num(post, r"主とする割合 ([\d.]+) は本解析の前に決めた値", 1)
    T = Table(
        "table8", {"ja": "表8", "en": "Table 8"},
        {"ja": f"拡張期の下降の扱いを変えたときの型3 の関連（探索・事後・{n_all} 名）{{posthoc}}{{src}}{{tier}}{{defs}}{{bold}}。",
         "en": f"Type-3 correlations when the handling of the diastolic decline is changed (exploratory, post hoc, {n_all} subjects){{posthoc}}{{src}}{{tier}}{{defs}}{{bold}}."},
        {"ja": [c["header_ja"] for c in t["columns"]],
         "en": ["Handling of the diastolic decline", "ΔT × PWV, tier A", "ΔT × PWV, tier C", "RI × resistance, tier A", "RI × resistance, tier C",
                "Pass rate", "ΔT offset from the supplied fiducial point (tier A)"]},
        notes=posthoc_notes("表6") + [
            ("src", source_note(t)),
            ("tier", {"ja": STAGE_NOTE["ja"] + "。通過率は凍結版と同じ収束検算（境界張り付き・成分の高さ・競合解の反射係数の差）を当てたときに採用になる割合、すなわち A 段に残る割合。",
                      "en": STAGE_NOTE["en"] + ". Pass rate: the fraction accepted under the same convergence checks as the frozen version "
                            "(parameters on a bound, component height, difference in reflection coefficient between competing solutions), i.e. the fraction remaining in tier A."}),
            ("defs", {"ja": "同じ 4,374 名の同じ拍に当てはめ直したもの。2 カーネル模型には拡張期の下降を受ける項が無いので、下降の扱いだけを変えた版を、"
                            "下降を説明する項を足す案と、下降を当てはめの対象から外す・残差の中で小さくする案に分けて並べる。"
                            f"主とする割合 {frac} は本解析の前に決めた値であり、これらの数値を見て選び直していない（lab_log 追記145）。"
                            "右端の列は第2成分のピークと同梱の拡張期側の特徴点との差の中央値（正は第2成分が遅い）。",
                      "en": f"Refits of the same beats of the same {n_all} subjects. The two-kernel model has no term for the diastolic decline, so versions "
                            "that change only its handling are grouped into those that add a term for the decline and those that exclude it from the fit "
                            f"or down-weight it in the residual. The main fraction {frac} was fixed before this analysis and was not re-chosen after seeing these "
                            "numbers (lab_log entry 145). The last column is the median difference between the peak of the second component and the "
                            "diastolic fiducial point supplied with PWDB (positive, second component later)."}),
            ("bold", {"ja": f"太字は規準（{N['meta']['criterion']['bold_rule']}）を満たす値。",
                      "en": f"Bold, criterion met (median |ρ| ≥ {N['meta']['criterion']['min_median_abs_rho']:.2f} with the predicted sign in every stratum)."}),
        ],
        landscape=True)
    seen_group = None
    for r in t["rows"]:
        g = r.get("group")
        if g and g != seen_group:
            gk = GROUP_KEY[g]
            T.group(label_cell(LABELS["variant_groups"][gk]["ja"], LABELS["variant_groups"][gk]["en"], bold=True))
            seen_group = g
        T.add(variant_label(r["id"], "ja", bold=r["label_bold"]),
              rho_cell(r["dt_pwv_A"]), rho_cell(r["dt_pwv_C"]), rho_cell(r["ri_pvr_A"]), rho_cell(r["ri_pvr_C"]),
              value_cell(r["pass_rate"]), value_cell(r["dt_offset_vs_landmark_ms_A"]))
    return T


# ---------------------------------------------------------------------------
# 表9  B 段と雑音（02 表6c）
# ---------------------------------------------------------------------------

def build_table9(N: dict) -> Table:
    t = N["tables"]["表6c"]
    pre = "\n".join(t["preamble"])
    m = re.search(r"（([\d,]+) 名。雑音 1% で (\d+) 名、2% で (\d+) 名）、C 段 ＝ 採否を無視した全員（([\d,]+) 名）", pre)
    if not m:
        raise ValueError("表6c の前書きから B 段の人数を取り出せない")
    b_def = N["meta"]["stage_definitions"]["B"]["text"]
    T = Table(
        "table9", {"ja": "表9", "en": "Table 9"},
        {"ja": "B 段と、雑音を足した実行（型3・探索・事後）{posthoc}{src}{tier}{noise}{ref}{bold}。",
         "en": "Tier B and runs with added noise (type 3; exploratory, post hoc){posthoc}{src}{tier}{noise}{ref}{bold}."},
        {"ja": [c["header_ja"] for c in t["columns"]],
         "en": ["Handling of the diastolic decline", "Tier", "ΔT × PWV, no noise", "1%", "2%", "RI × resistance, no noise", "1%", "2%",
                "Pass rate, none / 1% / 2%"]},
        notes=posthoc_notes("表6c") + [
            ("src", source_note(t)),
            ("tier", {"ja": STAGE_NOTE["ja"] + "。この表の B 段 ＝ " + b_def + "、C 段 ＝ 採否を無視した全員（" + m.group(4) + " 名）。",
                      "en": STAGE_NOTE["en"] + f". Tier B here: subjects accepted by all three of the frozen fit, the 0.65T truncation and the "
                            f"derivative-domain version ({m.group(1)}; {m.group(2)} at 1% noise, {m.group(3)} at 2%); tier C, all type-3 subjects ({m.group(4)})."}),
            ("noise", {"ja": "雑音は拍の峰から谷までの振幅に対する比を標準偏差とする白色ガウス雑音で、波形の型と特徴点は雑音を足す前の拍から付けた。",
                       "en": "Noise is white Gaussian noise whose SD is the stated fraction of the peak-to-trough amplitude of the beat; waveform type and "
                             "fiducial points were assigned on the beat before noise was added."}),
            ("ref", {"ja": "同梱の特徴点法の C 段は雑音を足していない拍の値であり、雑音の列では同じ条件の比較になっていない。",
                     "en": "The fiducial-point reference (tier C) is computed on the noise-free beats, so in the noise columns it is not a like-for-like comparison."}),
            ("bold", {"ja": f"太字は規準（{N['meta']['criterion']['bold_rule']}）を満たす値。",
                      "en": f"Bold, criterion met (median |ρ| ≥ {N['meta']['criterion']['min_median_abs_rho']:.2f} with the predicted sign in every stratum)."}),
        ],
        landscape=True)
    for r in t["rows"]:
        lab = variant_label(r["id"], "ja")
        if r.get("label_carried"):
            lab = Cell({"ja": "", "en": ""}, kind="label", csv_text={"ja": plain(lab.text["ja"]), "en": plain(lab.text["en"])})
        cells = [lab, Cell({"ja": r["stage"], "en": r["stage"]}, nums=[], kind="label")]
        for key in ("dt_pwv", "ri_pvr"):
            for nl in t["noise_levels"]:
                c = r[key][nl]
                cells.append(Cell({"ja": c["raw"], "en": en_from_raw(c["raw"])},
                                  nums=[] if c.get("same_as_left") else [c["rho"]],
                                  bold=c["bold"] and not c.get("same_as_left")))
        pr = r["pass_rate"]
        if pr:
            cells.append(Cell({"ja": pr["raw"], "en": pr["raw"].replace("／", " / ")}, nums=[pr[nl] for nl in t["noise_levels"]]))
        elif r.get("reference"):
            # 02 表6c: 参考行（特徴点法。当てはめが無いので通過率は定義されない）は「—」。JSON は空欄と区別しないので参考の印で補う
            cells.append(Cell({"ja": "—", "en": "—"}, nums=[]))
        else:
            cells.append(Cell({"ja": "", "en": ""}, nums=[]))
        T.add(*cells)
    return T


# ---------------------------------------------------------------------------
# 補足表 S1  端に達した母数（02 表6b）
# ---------------------------------------------------------------------------

def build_tableS1(N: dict) -> Table:
    t = N["tables"]["表6b"]
    T = Table(
        "tableS1", {"ja": "補足表 S1", "en": "Table S1"},
        {"ja": "当てはめが探索範囲の端に達した母数（型3・C 段・上位 2 つ）{posthoc}{src}{tier}{params}{bold}。",
         "en": "Parameters that reached a search bound (type 3, tier C, top two){posthoc}{src}{tier}{params}{bold}."},
        {"ja": [c["header_ja"] for c in t["columns"]],
         "en": ["Handling of the diastolic decline", "Fraction of beats at a bound", "Breakdown"]},
        notes=posthoc_notes("表6b") + [
            ("src", source_note(t)),
            ("tier", STAGE_NOTE),
            ("params", {"ja": re.sub(r"^.*?説明できる。", "", re.sub(r"\s*\n\s*", "", t["preamble"][0])),
                        "en": "a1, μ1, σ1 and α1 are the height, position, width and skewness of the first component; a2, Δμ, σ2 and α2 the same for the "
                              "second; g and τ the gain and time constant of the reservoir term. ':lo' lower bound, ':hi' upper bound. A beat can reach "
                              "bounds in several parameters, so the fractions do not sum to 1."}),
            ("bold", {"ja": "太字は縮退を示す母数（本文で論じる）。", "en": "Bold, the parameters discussed in the text as the sign of degeneracy."}),
        ])
    for r in t["rows"]:
        items_ja, items_en, par_ja, par_en, nums = [], [], [], [], []
        for it in r["breakdown"]["items"]:
            s = unbold(it["raw"])
            if it["bold"]:
                s = f"**{s}**"
            if it["parenthetical"]:          # 02 表6b: 「α1:hi 0.25・μ1:lo 0.02（τ:lo 0.00）」のように括弧書きは区切り無しで続く
                par_ja.append(f"（{s}）")
                par_en.append(f" ({s})")
            else:
                items_ja.append(s)
                items_en.append(s)
            nums.append(it["fraction"])
        T.add(variant_label(r["id"], "ja"), value_cell(r["boundary_fraction"]),
              Cell({"ja": "・".join(items_ja) + "".join(par_ja), "en": "; ".join(items_en) + "".join(par_en)}, nums=nums))
    return T


# ---------------------------------------------------------------------------
# 補足表 S2a・S2b  線形分離の後進波との照合と、圧波形に当てた結果（02 表7a・7b）
# ---------------------------------------------------------------------------

METHOD_EN = {"pda_ppg": "Decomposition (PPG)", "pda_pressure": "Decomposition (pressure)", "landmark": "Fiducial-point (supplied)"}
S2B_EN = {"ppg": "PPG (primary analysis of this paper)", "pressure": "Digital pressure waveform",
          "landmark": "Fiducial-point analysis (supplied; PPG; reference)"}


def build_tableS2(N: dict) -> list[Table]:
    ta, tb = N["tables"]["表7a"], N["tables"]["表7b"]
    n_all = fmt_int(N["meta"]["n_subjects_decision_test"])
    post_a = ta["postscript"][0].replace("\n", "")
    n_excl = find_num(post_a, r"後進波が実質的に無い ([\d,]+) 名を除く", 1)
    T = Table(
        "tableS2a", {"ja": "補足表 S2a", "en": "Table S2a"},
        {"ja": "手法の ΔT と線形分離の後進波の到達時間 ΔT_true との差（C 段）{posthoc}{src}{tier}{sep}{defs}。",
         "en": "Difference between each method's ΔT and the arrival time ΔT_true of the linearly separated backward wave (tier C){posthoc}{src}{tier}{sep}{defs}."},
        {"ja": [c["header_ja"] for c in ta["columns"]], "en": ["Waveform type", "Method", "n", "Median difference", "Pooled Spearman ρ"]},
        notes=posthoc_notes("表7a") + [
            ("src", source_note(ta)),
            ("tier", STAGE_NOTE),
            ("sep", {"ja": "PWDB の指尖の圧 P と流速 U から線形の波分離 P_f = (ΔP + ρc·ΔU)/2、P_b = (ΔP − ρc·ΔU)/2（ρc は収縮期初期の P–U ループの傾き、"
                           "基線は拍内の最小値）で後進波を作り、ΔT_true は P_b と P_f のピーク時刻の差。",
                     "en": "The backward wave was obtained from the digital pressure P and flow velocity U of PWDB by linear wave separation, "
                           "P_f = (ΔP + ρc·ΔU)/2 and P_b = (ΔP − ρc·ΔU)/2 (ρc, slope of the early-systolic P–U loop; baseline, minimum within the beat); "
                           "ΔT_true is the difference between the peak times of P_b and P_f."}),
            ("defs", {"ja": post_a,
                      "en": f"n, subjects in which the separation could be evaluated ({n_excl} excluded because ρc could not be estimated or the "
                            "backward wave was essentially absent). Difference = (method's ΔT) − (ΔT_true). ρ is the Spearman correlation pooled over age strata."}),
        ])
    for r in ta["rows"]:
        T.add(label_cell(r["type_ja"], f"Type {r['type_id'][-1]}", nums=[float(r["type_id"][-1])]),
              label_cell(r["method_ja"], METHOD_EN[r["method_id"]]),
              value_cell(r["n"]), value_cell(r["dt_diff_ms"]), value_cell(r["rho_pooled"]))

    post_b = tb["postscript"][0]
    mb = re.search(r"A 段 ([\d.]+)（(\d)/(\d)・成立）、C 段 ([\d.]+)（(\d)/(\d)・不成立）", post_b)
    if not mb:
        raise ValueError("表7b の注から型3 の値を取り出せない")
    Tb = Table(
        "tableS2b", {"ja": "補足表 S2b", "en": "Table S2b"},
        {"ja": f"同じ凍結版の分解を指尖の圧波形に当てた結果（年齢層内 Spearman・全例 {n_all} 名）{{posthoc}}{{src}}{{tier}}{{t3}}{{bold}}。",
         "en": f"The same frozen decomposition applied to the digital pressure waveform (within-stratum Spearman; all {n_all} subjects){{posthoc}}{{src}}{{tier}}{{t3}}{{bold}}."},
        {"ja": [c["header_ja"] for c in tb["columns"]],
         "en": ["Waveform fitted", "ΔT × PWV, tier A", "ΔT × PWV, tier C", "RI × resistance, tier A", "RI × resistance, tier C", "Acceptance rate"]},
        notes=posthoc_notes("表7b") + [
            ("src", source_note(tb)),
            ("tier", STAGE_NOTE),
            ("t3", {"ja": post_b,
                    "en": f"Restricted to type 3, ΔT of the pressure waveform is {mb.group(1)} ({mb.group(2)}/{mb.group(3)}) in tier A, meeting the criterion, "
                          f"and {mb.group(4)} ({mb.group(5)}/{mb.group(6)}) in tier C, not meeting it."}),
            ("bold", {"ja": f"太字は規準（{N['meta']['criterion']['bold_rule']}）を満たす値。",
                      "en": f"Bold, criterion met (median |ρ| ≥ {N['meta']['criterion']['min_median_abs_rho']:.2f} with the predicted sign in every stratum)."}),
        ])
    for r in tb["rows"]:
        Tb.add(label_cell(r["label_ja"], S2B_EN[r["id"]]),
               rho_cell(r["dt_pwv_A"]), rho_cell(r["dt_pwv_C"]), rho_cell(r["ri_pvr_A"]), rho_cell(r["ri_pvr_C"]),
               value_cell(r["adoption_rate"]))
    return [T, Tb]


def build_all(N: dict, P: dict) -> list[Table]:
    return ([build_table1(N, P), build_table2(N)] + build_table3(N) + [build_table4(N)] + build_table5(N)
            + build_table6(N) + [build_table7(N, P), build_table8(N), build_table9(N), build_tableS1(N)] + build_tableS2(N))


def load_prespec(path: Path = PRESPEC) -> dict:
    with open(path, encoding="utf-8") as f:
        return json.load(f)


# ---------------------------------------------------------------------------
# 書き出し: Markdown・CSV・docx・凡例
# ---------------------------------------------------------------------------

DOC_INTRO = {
    "ja": ("数値はすべて `data/paper2_numbers.json`（`data/extract_tables.py` が `../02_tables.md` を機械で読んだもの）から "
           "`build_tables.py` が取り、手で打っていない。判定（成立・不成立）は `docs/research/roadmap_v1.md` §9 のまま動かさない。"
           "表4 以降は決定試験の判定の後に回した探索・事後の記述であり、判定には用いない。段（A・B・C）を出す表には段の注釈を脚注に置く。"
           "脚注の印（a, b, …）は見出し文または列見出しに付けた。"),
    "en": ("All numbers are read by `build_tables.py` from `data/paper2_numbers.json` (a mechanical parse of `../02_tables.md` by "
           "`data/extract_tables.py`); none is typed by hand. Verdicts (pass/fail) are those of `docs/research/roadmap_v1.md` §9 and are not "
           "changed. Tables 4 onward are exploratory, post hoc descriptions made after the prespecified decision and are not used for it. "
           "Every table that shows tiers A/B/C carries the tier note as a footnote. Footnote letters (a, b, …) are attached to the caption or to column headers."),
}
DOC_TITLE = {"ja": "論文2 表集（和文）", "en": "Paper 2: tables (English)"}


def md_table_lines(T: Table, lang: str) -> list[str]:
    hdr = [T.mark(md_escape(h)) for h in T.headers[lang]]
    lines = ["| " + " | ".join(hdr) + " |", "|" + "---|" * len(hdr)]
    for r in T.rows:
        if isinstance(r, tuple):
            lines.append("| " + r[1].md(lang) + " |" + " |" * (len(hdr) - 1))
        else:
            lines.append("| " + " | ".join(T.mark(c.md(lang)) for c in r) + " |")
    return lines


def render_md(tables: list[Table], lang: str) -> str:
    out = [f"# {DOC_TITLE[lang]}", "", DOC_INTRO[lang], ""]
    for T in tables:
        out += [f"## {T.number[lang]}", "", f"**{T.number[lang]}.** {T.mark(T.caption[lang])}", ""]
        out += md_table_lines(T, lang)
        out.append("")
        for i, (_k, txt) in enumerate(T.notes):
            out += [f"{sup('abcdefghijklmnopqrstuvwxyz'[i])} {txt[lang]}", ""]
    return "\n".join(out).rstrip() + "\n"


def strip_marks(h: str) -> str:
    return re.sub(r"\{[a-z_0-9+]+\}", "", h)


def write_csv(T: Table, lang: str, path: Path) -> int:
    has_group = any(isinstance(r, tuple) for r in T.rows)
    header = [strip_marks(h) for h in T.headers[lang]] + (["group"] if has_group else [])
    n = 0
    with open(path, "w", encoding="utf-8", newline="") as f:
        w = csv.writer(f)
        w.writerow(header)
        group = ""
        for r in T.rows:
            if isinstance(r, tuple):
                group = plain(r[1].text[lang])
                continue
            row = [strip_marks(c.csv[lang] if c.csv else plain(c.text[lang])) for c in r]
            w.writerow(row + ([group] if has_group else []))
            n += 1
    return n


def legend_text(tables: list[Table], lang: str) -> str:
    out = [DOC_TITLE[lang], "", DOC_INTRO[lang].replace("`", ""), ""]
    for T in tables:
        out.append(f"{T.number[lang]}. {plain(strip_marks(T.caption[lang]))}")
        for i, (_k, txt) in enumerate(T.notes):
            out.append(f"  ({'abcdefghijklmnopqrstuvwxyz'[i]}) {txt[lang]}")
        out.append("")
    return "\n".join(out).rstrip() + "\n"


_TOKEN_RE = re.compile(r"(\*\*|<sup>[a-z](?:, [a-z])*</sup>)")


def write_docx(tables: list[Table], lang: str, path: Path) -> None:
    from docx import Document
    from docx.enum.section import WD_ORIENT, WD_SECTION
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Mm, Pt

    font_name = "游明朝" if lang == "ja" else "Times New Roman"
    body_pt = 10.5 if lang == "ja" else 10.0
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = font_name
    st.font.size = Pt(body_pt)
    rpr = st.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), font_name)

    def set_page(sec, landscape: bool):
        if landscape:
            sec.orientation = WD_ORIENT.LANDSCAPE
            sec.page_width, sec.page_height = Mm(297), Mm(210)
        else:
            sec.orientation = WD_ORIENT.PORTRAIT
            sec.page_width, sec.page_height = Mm(210), Mm(297)
        sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Mm(20)

    def style_run(run, size, bold=None, sup_=False):
        run.font.name = font_name
        run.font.size = Pt(size)
        r = run._element.get_or_add_rPr()
        rf = r.find(qn("w:rFonts"))
        if rf is None:
            rf = OxmlElement("w:rFonts")
            r.append(rf)
        rf.set(qn("w:eastAsia"), font_name)
        if bold is not None:
            run.bold = bold
        if sup_:
            run.font.superscript = True

    def add_marked(par, text: str, size: float, base_bold: bool = False):
        """`**` と <sup>x</sup> を解釈して run を足す。逆引用符は外す。"""
        text = text.replace("`", "").replace("\\|", "|")
        bold = base_bold
        for tok in _TOKEN_RE.split(text):
            if not tok:
                continue
            if tok == "**":
                bold = not bold
                continue
            m = re.fullmatch(r"<sup>([a-z](?:, [a-z])*)</sup>", tok)
            if m:
                style_run(par.add_run(m.group(1)), size, bold=False, sup_=True)
                continue
            style_run(par.add_run(tok), size, bold=bold)

    set_page(doc.sections[0], False)
    p = doc.add_paragraph()
    style_run(p.add_run(DOC_TITLE[lang]), body_pt + 3, bold=True)
    add_marked(doc.add_paragraph(), DOC_INTRO[lang], body_pt)
    cur_landscape = False
    for i, T in enumerate(tables):
        if T.landscape != cur_landscape:
            set_page(doc.add_section(WD_SECTION.NEW_PAGE), T.landscape)
            cur_landscape = T.landscape
        elif i > 0:
            doc.add_page_break()
        cell_pt = 9.0 if T.small else body_pt
        cap = doc.add_paragraph()
        style_run(cap.add_run(f"{T.number[lang]}. "), body_pt, bold=True)
        add_marked(cap, T.mark(T.caption[lang]), body_pt)
        ncols = T.ncols(lang)
        tbl = doc.add_table(rows=1, cols=ncols)
        tbl.style = "Table Grid"
        tbl.autofit = True
        hdr = tbl.rows[0]
        tr_pr = hdr._tr.get_or_add_trPr()
        th = OxmlElement("w:tblHeader")
        th.set(qn("w:val"), "true")
        tr_pr.append(th)
        for j, h in enumerate(T.headers[lang]):
            add_marked(hdr.cells[j].paragraphs[0], T.mark(h), cell_pt, base_bold=True)
        for r in T.rows:
            row = tbl.add_row()
            if isinstance(r, tuple):
                merged = row.cells[0].merge(row.cells[ncols - 1])
                add_marked(merged.paragraphs[0], unbold(r[1].text[lang]), cell_pt, base_bold=True)
                continue
            for j, c in enumerate(r):
                add_marked(row.cells[j].paragraphs[0], T.mark(c.md(lang)), cell_pt)
        for k, (_key, txt) in enumerate(T.notes):
            np_ = doc.add_paragraph()
            style_run(np_.add_run("abcdefghijklmnopqrstuvwxyz"[k]), body_pt, sup_=True)
            add_marked(np_, " " + txt[lang], body_pt)
    doc.save(str(path))


def write_all(tables: list[Table], lang: str, out_dir: Path, legend_dir: Path) -> list[Path]:
    out_dir.mkdir(parents=True, exist_ok=True)
    (out_dir / "csv").mkdir(exist_ok=True)
    legend_dir.mkdir(parents=True, exist_ok=True)
    paths = []
    p = out_dir / f"tables_{lang}.md"
    p.write_text(render_md(tables, lang), encoding="utf-8")
    paths.append(p)
    for T in tables:
        q = out_dir / "csv" / f"{T.csv_name}_{lang}.csv"
        write_csv(T, lang, q)
        paths.append(q)
    d = out_dir / f"tables_{lang}.docx"
    write_docx(tables, lang, d)
    paths.append(d)
    lp = legend_dir / f"tables_legend_{lang}.txt"
    lp.write_text(legend_text(tables, lang), encoding="utf-8")
    paths.append(lp)
    return paths


# ---------------------------------------------------------------------------
# 自己検査
# ---------------------------------------------------------------------------

def _banned_terms():
    """用語検査器（analysis/scripts/check_terminology.py）の禁止語の表を読む。写さない。"""
    path = REPO / "analysis" / "scripts" / "check_terminology.py"
    if not path.exists():
        return None, None
    spec = importlib.util.spec_from_file_location("check_terminology", path)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod.BANNED, mod.PRAGMA


def _banned_hits(banned, pragma, text: str) -> list[str]:
    hits = []
    for line in text.split("\n"):
        if pragma and pragma in line:
            continue
        for term, _alt, allow in banned:
            if term in line and not (allow and re.sub(allow, "", line).find(term) < 0):
                hits.append(f"{term}: {line[:50]}")
    return hits


def parse_md_tables(md: str) -> list[list[list[str]]]:
    blocks, cur = [], None
    for line in md.splitlines():
        if line.startswith("|"):
            cells = [c.strip() for c in re.split(r"(?<!\\)\|", line.strip())[1:-1]]
            cur = cur or []
            cur.append(cells)
        elif cur:
            blocks.append(cur)
            cur = None
    if cur:
        blocks.append(cur)
    return blocks


def docx_texts(path: Path) -> tuple[list[str], int, bool]:
    from docx import Document
    doc = Document(str(path))
    texts = [p.text for p in doc.paragraphs]
    header_bold = True
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                texts.append(c.text)
        for c in t.rows[0].cells:
            if not all(r.bold for p in c.paragraphs for r in p.runs if r.text.strip() and not r.font.superscript):
                header_bold = False
    return texts, len(doc.tables), header_bold


def selftest() -> int:
    import tempfile
    ok_all = True

    def rep(name, ok, detail=""):
        nonlocal ok_all
        ok_all &= bool(ok)
        print(f"  {'PASS' if ok else 'FAIL'}  {name}" + (f"  ({detail})" if detail and not ok else ""))

    banned, pragma = _banned_terms()
    N0, P0 = load_numbers(), load_prespec()
    tables = build_all(N0, P0)
    fresh = build_all(load_numbers(), load_prespec())     # JSON を読み直して組み直したもの（照合の相手）
    rep("表の数が 16（表1〜表9・S1・S2a・S2b）", len(tables) == 16, str(len(tables)))
    t1_verdict = {"凍結版": "pda_frozen", "第2版 歪みガウス": "pda_v2_skewgauss", "第2版 ガンマ": "pda_v2_gamma",
                  "特徴点法（Charlton": "landmark", "早期振幅比": "amp_ratio", "陽性対照": "ptt_control"}
    r1 = {r["id"]: r for r in N0["tables"]["表1"]["rows"]}
    with tempfile.TemporaryDirectory() as tmp:
        tmpd = Path(tmp)
        for lang in LANGS:
            print(f"[{lang}]")
            md = render_md(tables, lang)
            blocks = parse_md_tables(md)
            rep("Markdown の表の数が組んだ表と同じ", len(blocks) == len(tables), f"{len(blocks)} != {len(tables)}")
            shape_ok, num_bad, bold_bad, n_num = True, [], [], 0
            for T, Tf, blk in zip(tables, fresh, blocks):
                hdr, rows = blk[0], blk[2:]
                if len(hdr) != T.ncols(lang) or any(len(r) != T.ncols(lang) for r in rows) or len(rows) != len(T.rows):
                    shape_ok = False
                    continue
                for r_model, r_fresh, r_md in zip(T.rows, Tf.rows, rows):
                    if isinstance(r_model, tuple):
                        if not (r_md[0].startswith("**") and all(c == "" for c in r_md[1:])):
                            bold_bad.append(f"{T.tid}: group row {r_md[0][:20]}")
                        continue
                    for c_model, c_fresh, s in zip(r_model, r_fresh, r_md):
                        s = _SUP_TAG_RE.sub("", s)
                        if c_fresh.nums is not None:
                            n_num += 1
                            if not same_nums(nums_in(s), c_fresh.nums):
                                num_bad.append(f"{T.tid}: {s[:30]} -> {nums_in(s)} != {c_fresh.nums}")
                        if c_fresh.bold and not (s.startswith("**") and s.endswith("**")):
                            bold_bad.append(f"{T.tid}: {s[:30]} should be bold")
                        if not c_fresh.bold and "**" not in c_fresh.text[lang] and "**" in s:
                            bold_bad.append(f"{T.tid}: {s[:30]} should not be bold")
            rep("見出しとセルの数が各行でそろい、行数も一致する（Markdown が描ける）", shape_ok)
            rep(f"数値セルの数値が JSON を読み直した値と一致（{n_num} セル）", not num_bad, "; ".join(num_bad[:4]))
            rep("太字が JSON の印どおり（群の行も太字）", not bold_bad, "; ".join(bold_bad[:4]))
            # 表1 の判定が 02_tables 表1 と同じ
            T1 = tables[0]
            v_ok = True
            for key, rid in t1_verdict.items():
                row = next(r for r in T1.rows if key in r[0].text["ja"])
                want = r1[rid]["verdict"]["verdict"]
                got = row[4].text[lang]
                if not got.startswith(want if lang == "ja" else VERDICT_EN[want]):
                    v_ok = False
            rep("表1 の決定試験 6 行の判定が 02_tables 表1 の判定と同じ", v_ok)
            # 文章の列の数値が出典の文字列に含まれる（表1 規則・出典、表5b、表7）
            sub_bad = []
            for T in tables:
                for r in T.data_rows():
                    for c in r:
                        if c.kind == "text" and c.src:
                            have = {round(x, 6) for x in nums_in(c.src)}
                            got = {round(x, 6) for x in nums_in(plain(c.text[lang]))}
                            if not got <= have:
                                sub_bad.append(f"{T.tid}: {sorted(got - have)} in {plain(c.text[lang])[:30]}")
            rep("文章の列の数値が出典の文字列（JSON）に含まれる（表1・表5b・表7）", not sub_bad, "; ".join(sub_bad[:4]))
            # 表7 の長さ
            T7 = next(T for T in tables if T.tid == "table7")
            long_ = []
            for r in T7.data_rows():
                for c in r:
                    if c.kind == "text":
                        s = plain(c.text[lang])
                        if lang == "ja" and len(s) > MAX_JA_CHARS_T7:
                            long_.append(f"{len(s)}: {s[:20]}")
                        if lang == "en" and len(s.split()) > MAX_EN_WORDS_T7:
                            long_.append(f"{len(s.split())}: {s[:20]}")
            rep(f"表7 の文章の列が上限以内（和 {MAX_JA_CHARS_T7} 字・英 {MAX_EN_WORDS_T7} 語）", not long_, "; ".join(long_[:4]))
            # 段の注釈・探索事後の注記
            tier_word = "段" if lang == "ja" else "ier "
            miss_tier, miss_post = [], []
            for i, T in enumerate(tables):
                notes = " ".join(txt[lang] for _k, txt in T.notes)
                shows_tier = any(tier_word in strip_marks(h) for h in T.headers[lang]) or any(
                    tier_word in c.text[lang] for r in T.data_rows() for c in r) or (tier_word in T.caption[lang])
                if shows_tier and STAGE_NOTE[lang] not in notes:
                    miss_tier.append(T.tid)
                # 探索・事後の注記（labels.json）は「判定には用いない」を含む
                if i >= 5 and not (POSTHOC_NOTE[lang] in notes and
                                   ("判定には用いない" in notes if lang == "ja" else "not used for the" in notes)):
                    miss_post.append(T.tid)
            rep("段（A・B・C）を出す表に段の注釈がある", not miss_tier, str(miss_tier))
            rep("表4 以降に探索・事後の注記と「判定には用いない」がある", not miss_post, str(miss_post))
            if lang == "en":
                # 英文の地の文に特徴点の旧い呼び名（_LEGACY_EN_TERM）を使わない（fiducial point と書く）。
                # 逆引用符の中の台本名・関数名は除く
                prose = re.sub(r"`[^`]*`", "", md + "\n" + legend_text(tables, lang))
                lm = re.findall("(?i)" + _LEGACY_EN_TERM, prose)
                rep(f"英文の地の文に {_LEGACY_EN_TERM} が無い（fiducial point を使う）", not lm, f"{len(lm)} 件")
            # 書き出して読み直す
            paths = write_all(tables, lang, tmpd, tmpd / "out")
            texts, n_tables, hdr_bold = docx_texts(tmpd / f"tables_{lang}.docx")
            rep("docx が開けて表の数が 16", n_tables == len(tables), str(n_tables))
            rep("docx の見出し行が太字", hdr_bold)
            csv_bad = []
            for T in tables:
                with open(tmpd / "csv" / f"{T.csv_name}_{lang}.csv", encoding="utf-8", newline="") as f:
                    rows = list(csv.reader(f))
                if len(rows) - 1 != len(T.data_rows()) or any(len(r) != len(rows[0]) for r in rows):
                    csv_bad.append(f"{T.tid}: {len(rows) - 1} != {len(T.data_rows())}")
            rep("csv の行数が表と一致し、列数がそろう", not csv_bad, "; ".join(csv_bad[:4]))
            leg = legend_text(tables, lang)
            if banned is None:
                print("  （用語検査器が無いので禁止語の確認は飛ばした）")
            else:
                hits = _banned_hits(banned, pragma, md + "\n" + leg + "\n" + "\n".join(texts))
                rep("禁止語が無い（md・凡例・docx。検査器の表）", not hits, "; ".join(hits[:4]))
    chk = REPO / "analysis" / "scripts" / "check_terminology.py"
    if chk.exists():
        r = subprocess.run([sys.executable, str(chk), str(HERE.relative_to(REPO))], cwd=str(REPO), capture_output=True, text=True)
        rep("リポジトリの用語検査（figtab）が通る", r.returncode == 0, r.stdout.strip().splitlines()[-1] if r.stdout.strip() else "")
    print("RESULT", "PASS" if ok_all else "FAIL")
    return 0 if ok_all else 1


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--lang", choices=["ja", "en"], default=None, help="省略すると両方")
    ap.add_argument("--selftest", action="store_true")
    ap.add_argument("--out-dir", default=str(HERE), help="tables_*.md・csv/・tables_*.docx を書く場所（既定はこの台本の場所）")
    ap.add_argument("--legend-dir", default=str(OUT_LEGEND), help="凡例文 tables_legend_*.txt を書く場所（既定は out/）")
    a = ap.parse_args()
    if a.selftest:
        return selftest()
    tables = build_all(load_numbers(), load_prespec())
    for lang in ([a.lang] if a.lang else list(LANGS)):
        for p in write_all(tables, lang, Path(a.out_dir), Path(a.legend_dir)):
            print(p)
    return 0


if __name__ == "__main__":
    sys.exit(main())
